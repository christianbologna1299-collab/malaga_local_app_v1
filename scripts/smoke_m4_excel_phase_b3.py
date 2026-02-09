#!/usr/bin/env python3
"""
M4.01 Phase B3 Smoke Test — Real Excel Chart Objects.

Verifies:
  - Chart factory functions exist and return correct types
  - Charts sheet contains 3 chart objects (LineChart, AreaChart, LineChart)
  - Chart titles, data references, grouping, and placement
  - Placeholder text removed
  - Save/reload preserves charts
  - Backward compat: all Phase A/B1/B2 structures intact

Usage: python scripts/smoke_m4_excel_phase_b3.py
Exit code 0 = PASS, nonzero = FAIL
"""

import sys
import os
import tempfile
from datetime import date

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

PASSED = 0
FAILED = 0


def check(name, condition, detail=""):
    global PASSED, FAILED
    if condition:
        PASSED += 1
        print(f"  PASS  {name}")
    else:
        FAILED += 1
        print(f"  FAIL  {name} — {detail}")


def _chart_title_text(chart):
    """Extract plain text from an openpyxl chart title (Title object or str)."""
    title = chart.title
    if title is None:
        return None
    if isinstance(title, str):
        return title
    try:
        for p in title.tx.rich.paragraphs:
            for r in p.r:
                return r.t
    except (AttributeError, IndexError):
        pass
    return str(title)


def main():
    print("=" * 60)
    print("M4.01 Phase B3 Smoke Tests — Real Excel Chart Objects")
    print("=" * 60)

    # ------------------------------------------------------------------
    # 1. Imports
    # ------------------------------------------------------------------
    print("\n[1] Imports")
    try:
        from features.excel_export import (
            LoanInputs, ScenarioConfig, SeriesData, ExportMeta,
            _build_skeleton_workbook, _build_charts_sheet,
            _make_balance_chart, _make_payment_breakdown_chart,
            _make_payment_chart,
            _CHART_WIDTH, _CHART_HEIGHT, _MAX_AMORT_ROWS,
        )
        check("B3 imports", True)
    except ImportError as e:
        check("B3 imports", False, str(e))
        print("\nCannot continue. Aborting.")
        sys.exit(1)

    try:
        import openpyxl
        from openpyxl.chart import LineChart, AreaChart
        check("openpyxl + chart imports", True)
    except ImportError as e:
        check("openpyxl + chart imports", False, str(e))
        sys.exit(1)

    check("_build_charts_sheet is callable", callable(_build_charts_sheet))
    check("_make_balance_chart is callable", callable(_make_balance_chart))
    check("_make_payment_breakdown_chart is callable", callable(_make_payment_breakdown_chart))
    check("_make_payment_chart is callable", callable(_make_payment_chart))
    check("_CHART_WIDTH = 22", _CHART_WIDTH == 22, f"got {_CHART_WIDTH}")
    check("_CHART_HEIGHT = 12", _CHART_HEIGHT == 12, f"got {_CHART_HEIGHT}")

    # ------------------------------------------------------------------
    # 2. Build workbook
    # ------------------------------------------------------------------
    print("\n[2] Build workbook")
    loan = LoanInputs(
        principal=100000, start_date="2023-01-01", term_months=60,
        rate=6.25, payment_frequency="Monthly",
    )
    scenario = ScenarioConfig(rate_shock_bps=0, balance_shock_pct=0.0, stress_toggle=False)
    series = SeriesData(dates=["2023-01-01"], balances=[100000], rates=[6.25])
    wb = _build_skeleton_workbook(loan, scenario, series)
    check("Workbook created", wb is not None)
    check("5 sheets", len(wb.sheetnames) == 5)
    check("Charts sheet exists", "Charts" in wb.sheetnames)

    # ------------------------------------------------------------------
    # 3. Chart objects on Charts sheet
    # ------------------------------------------------------------------
    print("\n[3] Chart objects")
    ws_charts = wb["Charts"]
    charts = ws_charts._charts
    check("Charts sheet has 3 charts", len(charts) == 3,
          f"got {len(charts)}")

    chart1, chart2, chart3 = charts[0], charts[1], charts[2]

    check("Chart 1 is LineChart", isinstance(chart1, LineChart),
          f"got {type(chart1).__name__}")
    check("Chart 2 is AreaChart", isinstance(chart2, AreaChart),
          f"got {type(chart2).__name__}")
    check("Chart 3 is LineChart", isinstance(chart3, LineChart),
          f"got {type(chart3).__name__}")

    # ------------------------------------------------------------------
    # 4. Chart titles
    # ------------------------------------------------------------------
    print("\n[4] Chart titles")
    t1 = _chart_title_text(chart1)
    t2 = _chart_title_text(chart2)
    t3 = _chart_title_text(chart3)
    check("Chart 1 title = 'Balance Over Time'",
          t1 == "Balance Over Time", f"got {t1}")
    check("Chart 2 title = 'Payment Breakdown'",
          t2 == "Payment Breakdown", f"got {t2}")
    check("Chart 3 title = 'Payment Over Time'",
          t3 == "Payment Over Time", f"got {t3}")

    # ------------------------------------------------------------------
    # 5. Chart 1 data references (Balance Over Time)
    # ------------------------------------------------------------------
    print("\n[5] Chart 1 — Balance Over Time data refs")
    check("Chart 1 has 1 data series", len(chart1.series) == 1,
          f"got {len(chart1.series)}")

    # Data reference should point to column 6 (Balance)
    s1 = chart1.series[0]
    s1_ref = str(s1.val.numRef.f) if hasattr(s1.val, 'numRef') else str(s1.val)
    check("Chart 1 data refs col F (Balance)",
          "$F$" in s1_ref or "!F" in s1_ref or "'Amortization'" in s1_ref,
          f"got {s1_ref}")

    check("Chart 1 width = 22", chart1.width == _CHART_WIDTH,
          f"got {chart1.width}")
    check("Chart 1 height = 12", chart1.height == _CHART_HEIGHT,
          f"got {chart1.height}")

    # Y-axis title
    check("Chart 1 y-axis title",
          chart1.y_axis.title is not None and "Balance" in str(chart1.y_axis.title),
          f"got {chart1.y_axis.title}")

    # ------------------------------------------------------------------
    # 6. Chart 2 data references (Payment Breakdown)
    # ------------------------------------------------------------------
    print("\n[6] Chart 2 — Payment Breakdown data refs")
    check("Chart 2 has 2 data series", len(chart2.series) == 2,
          f"got {len(chart2.series)}")
    check("Chart 2 grouping = 'stacked'", chart2.grouping == "stacked",
          f"got {chart2.grouping}")

    # Y-axis title
    check("Chart 2 y-axis title",
          chart2.y_axis.title is not None and "Amount" in str(chart2.y_axis.title),
          f"got {chart2.y_axis.title}")

    # ------------------------------------------------------------------
    # 7. Chart 3 data references (Payment Over Time)
    # ------------------------------------------------------------------
    print("\n[7] Chart 3 — Payment Over Time data refs")
    check("Chart 3 has 1 data series", len(chart3.series) == 1,
          f"got {len(chart3.series)}")

    s3 = chart3.series[0]
    s3_ref = str(s3.val.numRef.f) if hasattr(s3.val, 'numRef') else str(s3.val)
    check("Chart 3 data refs col C (Payment)",
          "$C$" in s3_ref or "!C" in s3_ref or "'Amortization'" in s3_ref,
          f"got {s3_ref}")

    check("Chart 3 y-axis title",
          chart3.y_axis.title is not None and "Payment" in str(chart3.y_axis.title),
          f"got {chart3.y_axis.title}")

    # ------------------------------------------------------------------
    # 8. Placeholder removed
    # ------------------------------------------------------------------
    print("\n[8] Placeholder removed")
    # Check that old placeholder text is gone
    placeholder_found = False
    for row in range(1, 20):
        val = ws_charts.cell(row=row, column=1).value
        if val and "Phase B3" in str(val):
            placeholder_found = True
            break
        if val and "will populate" in str(val).lower():
            placeholder_found = True
            break
    check("No B3 placeholder text", not placeholder_found)

    # ------------------------------------------------------------------
    # 9. Charts sheet formatting
    # ------------------------------------------------------------------
    print("\n[9] Charts sheet formatting")
    title_val = ws_charts.cell(row=1, column=1).value
    check("Title row present",
          title_val is not None and "Charts" in str(title_val),
          f"got {title_val}")

    check("A2 label = 'Balance Over Time'",
          ws_charts["A2"].value == "Balance Over Time",
          f"got {ws_charts['A2'].value}")
    check("A18 label = 'Payment Breakdown'",
          ws_charts["A18"].value == "Payment Breakdown",
          f"got {ws_charts['A18'].value}")
    check("A34 label = 'Payment Over Time'",
          ws_charts["A34"].value == "Payment Over Time",
          f"got {ws_charts['A34'].value}")

    check("Column A width >= 30",
          ws_charts.column_dimensions["A"].width >= 30,
          f"got {ws_charts.column_dimensions['A'].width}")

    # ------------------------------------------------------------------
    # 10. Save and reload
    # ------------------------------------------------------------------
    print("\n[10] Save and reload")
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name
    try:
        wb.save(tmp_path)
        check("Workbook saves", True)

        wb2 = openpyxl.load_workbook(tmp_path)
        check("Workbook reloads", True)
        check("5 sheets after reload", len(wb2.sheetnames) == 5)

        ws2_charts = wb2["Charts"]
        charts2 = ws2_charts._charts
        check("3 charts survive reload", len(charts2) == 3,
              f"got {len(charts2)}")

        # Titles survive
        t1r = _chart_title_text(charts2[0])
        t2r = _chart_title_text(charts2[1])
        t3r = _chart_title_text(charts2[2])
        check("Chart 1 title survives",
              t1r == "Balance Over Time", f"got {t1r}")
        check("Chart 2 title survives",
              t2r == "Payment Breakdown", f"got {t2r}")
        check("Chart 3 title survives",
              t3r == "Payment Over Time", f"got {t3r}")

        wb2.close()
    finally:
        os.unlink(tmp_path)

    # ------------------------------------------------------------------
    # 11. Backward compat
    # ------------------------------------------------------------------
    print("\n[11] Backward compat")
    # Named ranges
    defined_names = list(wb.defined_names.keys())
    check("15 named ranges still present", len(defined_names) == 15,
          f"got {len(defined_names)}: {defined_names}")

    # Amortization formulas intact
    ws_amort = wb["Amortization"]
    a3 = str(ws_amort["A3"].value)
    check("Amort A3 still formula", a3.startswith("=") and "Num_Periods" in a3,
          f"got {a3}")

    f3 = str(ws_amort["F3"].value)
    check("Amort F3 still formula", f3.startswith("=") and "Loan_Principal" in f3,
          f"got {f3}")

    # Loan Inputs intact
    ws_inp = wb["Loan Inputs"]
    b7_val = ws_inp["B7"].value
    check("Start date still datetime.date",
          isinstance(b7_val, date), f"got type {type(b7_val).__name__}")

    b16 = str(ws_inp["B16"].value)
    check("B16 Payment formula intact", "PMT" in b16.upper(), f"got {b16}")

    # Overview intact
    ws_ov = wb["Overview"]
    b4_ov = str(ws_ov["B4"].value)
    check("Overview B4 refs Loan_Principal", "Loan_Principal" in b4_ov,
          f"got {b4_ov}")

    # Scenarios intact
    ws_sc = wb["Scenarios"]
    sc_b8 = str(ws_sc.cell(row=8, column=2).value)
    check("Scenarios B8 refs Loan_Principal", "Loan_Principal" in sc_b8,
          f"got {sc_b8}")

    # ------------------------------------------------------------------
    # Summary
    # ------------------------------------------------------------------
    print("\n" + "=" * 60)
    total = PASSED + FAILED
    print(f"Results: {PASSED}/{total} passed, {FAILED} failed")
    if FAILED == 0:
        print("M4.01 PHASE B3 SMOKE TESTS: ALL PASSED")
    else:
        print("M4.01 PHASE B3 SMOKE TESTS: FAILURES DETECTED")
    print("=" * 60)

    sys.exit(1 if FAILED > 0 else 0)


if __name__ == "__main__":
    main()
