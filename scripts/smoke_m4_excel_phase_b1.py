#!/usr/bin/env python3
"""
M4.01 Phase B1 Smoke Test — Formula-Driven Amortization.

Verifies:
  - Derived named ranges (Periods_Per_Year, Periodic_Rate, Num_Periods, Payment)
  - Amort_Data named range
  - Amortization cells contain formulas (not static values)
  - IF guards reference Num_Periods
  - Period 1 formulas reference Loan_Principal / Periodic_Rate / Payment
  - General row formulas reference previous row (F{r-1})
  - Zebra striping on even-period rows
  - 360-row formula grid
  - Save/reload preserves formulas
  - Start date stored as datetime.date

Usage: python scripts/smoke_m4_excel_phase_b1.py
Exit code 0 = PASS, nonzero = FAIL
"""

import sys
import os
import tempfile
from datetime import date

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

PASSED = 0
FAILED = 0


def check(name, condition, detail=""):
    global PASSED, FAILED
    if condition:
        PASSED += 1
        print(f"  PASS  {name}")
    else:
        FAILED += 1
        print(f"  FAIL  {name} — {detail}")


def main():
    print("=" * 60)
    print("M4.01 Phase B1 Smoke Tests — Formula-Driven Amortization")
    print("=" * 60)

    # ------------------------------------------------------------------
    # 1. Imports
    # ------------------------------------------------------------------
    print("\n[1] Imports")
    try:
        from features.excel_export import (
            LoanInputs,
            ScenarioConfig,
            SeriesData,
            _build_skeleton_workbook,
            _MAX_AMORT_ROWS,
        )
        check("excel_export B1 imports", True)
    except ImportError as e:
        check("excel_export B1 imports", False, str(e))
        print("\nCannot continue. Aborting.")
        sys.exit(1)

    try:
        import openpyxl
        check("openpyxl importable", True)
    except ImportError as e:
        check("openpyxl importable", False, str(e))
        sys.exit(1)

    check("_MAX_AMORT_ROWS = 360", _MAX_AMORT_ROWS == 360,
          f"got {_MAX_AMORT_ROWS}")

    # ------------------------------------------------------------------
    # 2. Build workbook
    # ------------------------------------------------------------------
    print("\n[2] Build workbook")
    loan = LoanInputs(
        principal=100000,
        start_date="2023-01-01",
        term_months=60,
        rate=6.25,
        payment_frequency="Monthly",
    )
    scenario = ScenarioConfig(rate_shock_bps=0, balance_shock_pct=0.0, stress_toggle=False)
    series = SeriesData(
        dates=["2023-01-01"],
        balances=[100000],
        rates=[6.25],
    )
    wb = _build_skeleton_workbook(loan, scenario, series)
    check("Workbook created", wb is not None)

    # ------------------------------------------------------------------
    # 3. Derived named ranges
    # ------------------------------------------------------------------
    print("\n[3] Derived named ranges (B1)")
    defined_names = list(wb.defined_names.keys())
    b1_names = ["Periods_Per_Year", "Periodic_Rate", "Num_Periods", "Payment", "Amort_Data"]
    for name in b1_names:
        check(f"Named range '{name}' exists", name in defined_names,
              f"defined names: {defined_names}")

    # Phase A names still present
    phase_a_names = ["Loan_Principal", "Loan_Rate", "Loan_Term_Months",
                     "Loan_Start_Date", "Loan_Frequency"]
    for name in phase_a_names:
        check(f"Phase A '{name}' preserved", name in defined_names)

    check("Total named ranges >= 10",
          len(defined_names) >= 10,
          f"got {len(defined_names)}: {defined_names}")

    # ------------------------------------------------------------------
    # 4. Derived value formulas on Loan Inputs
    # ------------------------------------------------------------------
    print("\n[4] Derived value formulas (Loan Inputs)")
    ws_inputs = wb["Loan Inputs"]

    # Subheader row 12
    check("Row 12 col A = 'Derived Values'",
          ws_inputs.cell(row=12, column=1).value == "Derived Values",
          f"got {ws_inputs.cell(row=12, column=1).value}")

    # B13: Periods_Per_Year formula
    b13 = str(ws_inputs["B13"].value)
    check("B13 is formula", b13.startswith("="), f"got {b13}")
    check("B13 references Loan_Frequency", "Loan_Frequency" in b13, f"got {b13}")
    check("B13 handles Monthly", '"Monthly"' in b13, f"got {b13}")
    check("B13 handles Quarterly", '"Quarterly"' in b13, f"got {b13}")

    # B14: Periodic_Rate formula
    b14 = str(ws_inputs["B14"].value)
    check("B14 is formula", b14.startswith("="), f"got {b14}")
    check("B14 references Loan_Rate", "Loan_Rate" in b14, f"got {b14}")
    check("B14 references Periods_Per_Year", "Periods_Per_Year" in b14, f"got {b14}")

    # B15: Num_Periods formula
    b15 = str(ws_inputs["B15"].value)
    check("B15 is formula", b15.startswith("="), f"got {b15}")
    check("B15 references Loan_Term_Months", "Loan_Term_Months" in b15, f"got {b15}")

    # B16: Payment formula
    b16 = str(ws_inputs["B16"].value)
    check("B16 is formula", b16.startswith("="), f"got {b16}")
    check("B16 uses PMT function", "PMT" in b16.upper(), f"got {b16}")
    check("B16 references Loan_Principal", "Loan_Principal" in b16, f"got {b16}")

    # ------------------------------------------------------------------
    # 5. Start date stored as datetime.date
    # ------------------------------------------------------------------
    print("\n[5] Start date")
    b7 = ws_inputs["B7"].value
    check("B7 is datetime.date", isinstance(b7, date), f"got type {type(b7).__name__}")
    check("B7 = 2023-01-01", b7 == date(2023, 1, 1), f"got {b7}")

    # ------------------------------------------------------------------
    # 6. Amortization column structure
    # ------------------------------------------------------------------
    print("\n[6] Amortization structure")
    ws_amort = wb["Amortization"]

    # 6 column headers
    headers = [ws_amort.cell(row=2, column=c).value for c in range(1, 7)]
    check("6 column headers",
          headers == ["Period", "Date", "Payment", "Interest", "Principal", "Balance"],
          f"got {headers}")

    # ------------------------------------------------------------------
    # 7. Row 3 formulas (Period 1 — special case)
    # ------------------------------------------------------------------
    print("\n[7] Row 3 formulas (Period 1)")
    a3 = str(ws_amort["A3"].value)
    check("A3 is formula", a3.startswith("="), f"got {a3}")
    check("A3 IF guard references Num_Periods", "Num_Periods" in a3, f"got {a3}")

    b3 = str(ws_amort["B3"].value)
    check("B3 references Loan_Start_Date", "Loan_Start_Date" in b3, f"got {b3}")

    c3 = str(ws_amort["C3"].value)
    check("C3 references Payment", "Payment" in c3, f"got {c3}")

    d3 = str(ws_amort["D3"].value)
    check("D3 references Loan_Principal", "Loan_Principal" in d3, f"got {d3}")
    check("D3 references Periodic_Rate", "Periodic_Rate" in d3, f"got {d3}")

    e3 = str(ws_amort["E3"].value)
    check("E3 = C3-D3 (principal = payment - interest)", "C3" in e3 and "D3" in e3,
          f"got {e3}")

    f3 = str(ws_amort["F3"].value)
    check("F3 references Loan_Principal", "Loan_Principal" in f3, f"got {f3}")

    # ------------------------------------------------------------------
    # 8. Row 4+ formulas (general case)
    # ------------------------------------------------------------------
    print("\n[8] Row 4 formulas (general case)")
    a4 = str(ws_amort["A4"].value)
    check("A4 is formula with IF guard", a4.startswith("=") and "Num_Periods" in a4,
          f"got {a4}")

    b4 = str(ws_amort["B4"].value)
    check("B4 uses EDATE", "EDATE" in b4.upper(), f"got {b4}")
    check("B4 references B3 (prev date)", "B3" in b4, f"got {b4}")

    d4 = str(ws_amort["D4"].value)
    check("D4 references F3 (prev balance)", "F3" in d4, f"got {d4}")
    check("D4 references Periodic_Rate", "Periodic_Rate" in d4, f"got {d4}")

    f4 = str(ws_amort["F4"].value)
    check("F4 references F3 (prev balance)", "F3" in f4, f"got {f4}")

    # ------------------------------------------------------------------
    # 9. Last formula row (row 362 = period 360)
    # ------------------------------------------------------------------
    print("\n[9] Last formula row (row 362)")
    last_row = 3 + _MAX_AMORT_ROWS - 1  # 362
    a_last = str(ws_amort.cell(row=last_row, column=1).value)
    check(f"Row {last_row} A has formula", a_last.startswith("="), f"got {a_last}")
    check(f"Row {last_row} IF guard", "Num_Periods" in a_last, f"got {a_last}")

    f_last = str(ws_amort.cell(row=last_row, column=6).value)
    check(f"Row {last_row} F has formula", f_last.startswith("="), f"got {f_last}")
    check(f"Row {last_row} F refs prev row",
          f"F{last_row - 1}" in f_last, f"got {f_last}")

    # Row 363 should be empty (no formula beyond 360 rows)
    beyond_row = last_row + 1
    check(f"Row {beyond_row} A is empty",
          ws_amort.cell(row=beyond_row, column=1).value is None,
          f"got {ws_amort.cell(row=beyond_row, column=1).value}")

    # ------------------------------------------------------------------
    # 10. Zebra striping
    # ------------------------------------------------------------------
    print("\n[10] Zebra striping")
    # Row 3 (period 1, odd) — no zebra fill
    row3_fill = ws_amort.cell(row=3, column=1).fill.start_color.rgb
    # Row 4 (period 2, even) — should have zebra fill
    row4_fill = ws_amort.cell(row=4, column=1).fill.start_color.rgb

    check("Row 3 and Row 4 have different fills",
          row3_fill != row4_fill,
          f"row3={row3_fill}, row4={row4_fill}")

    # Check that the zebra fill is the expected light blue-gray
    check("Row 4 fill contains F2F6FA",
          "F2F6FA" in str(row4_fill).upper(),
          f"got {row4_fill}")

    # Row 5 (period 3, odd) — should match row 3
    row5_fill = ws_amort.cell(row=5, column=1).fill.start_color.rgb
    check("Row 5 (odd) matches row 3 (no zebra)",
          row5_fill == row3_fill,
          f"row3={row3_fill}, row5={row5_fill}")

    # ------------------------------------------------------------------
    # 11. Number formats
    # ------------------------------------------------------------------
    print("\n[11] Number formats")
    check("C3 (Payment) has currency format",
          ws_amort["C3"].number_format == '#,##0.00',
          f"got {ws_amort['C3'].number_format}")
    check("D3 (Interest) has currency format",
          ws_amort["D3"].number_format == '#,##0.00',
          f"got {ws_amort['D3'].number_format}")
    check("F3 (Balance) has currency format",
          ws_amort["F3"].number_format == '#,##0.00',
          f"got {ws_amort['F3'].number_format}")
    check("B3 (Date) has date format",
          ws_amort["B3"].number_format == 'YYYY-MM-DD',
          f"got {ws_amort['B3'].number_format}")

    # Derived value formats on Loan Inputs
    check("B14 (Periodic Rate) has precise format",
          ws_inputs["B14"].number_format == '0.000000%',
          f"got {ws_inputs['B14'].number_format}")
    check("B16 (Payment) has currency format",
          ws_inputs["B16"].number_format == '#,##0.00',
          f"got {ws_inputs['B16'].number_format}")

    # ------------------------------------------------------------------
    # 12. Save and reload (formulas preserved)
    # ------------------------------------------------------------------
    print("\n[12] Save and reload")
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name
    try:
        wb.save(tmp_path)
        check("Workbook saves without error", True)

        wb2 = openpyxl.load_workbook(tmp_path)
        check("Workbook reloads", True)

        # Formulas survive
        ws2_amort = wb2["Amortization"]
        a3_reload = str(ws2_amort["A3"].value)
        check("A3 formula survives reload",
              a3_reload.startswith("=") and "Num_Periods" in a3_reload,
              f"got {a3_reload}")

        f3_reload = str(ws2_amort["F3"].value)
        check("F3 formula survives reload",
              f3_reload.startswith("=") and "Loan_Principal" in f3_reload,
              f"got {f3_reload}")

        # Derived names survive
        reload_names = list(wb2.defined_names.keys())
        for name in b1_names:
            check(f"'{name}' survives reload", name in reload_names)

        # Start date survives as datetime
        ws2_inputs = wb2["Loan Inputs"]
        b7_reload = ws2_inputs["B7"].value
        check("B7 date survives reload",
              hasattr(b7_reload, 'year') and b7_reload.year == 2023,
              f"got {b7_reload} (type={type(b7_reload).__name__})")

        # Zebra survives
        r4_fill = ws2_amort.cell(row=4, column=1).fill.start_color.rgb
        check("Zebra fill survives reload",
              "F2F6FA" in str(r4_fill).upper(),
              f"got {r4_fill}")

        wb2.close()
    finally:
        os.unlink(tmp_path)

    # ------------------------------------------------------------------
    # 13. Amort_Data named range reference
    # ------------------------------------------------------------------
    print("\n[13] Amort_Data reference")
    amort_dn = wb.defined_names["Amort_Data"]
    amort_ref = str(amort_dn.attr_text)
    check("Amort_Data references Amortization sheet",
          "Amortization" in amort_ref, f"got {amort_ref}")
    check("Amort_Data spans A2:F362",
          "A$2" in amort_ref and "F$362" in amort_ref,
          f"got {amort_ref}")

    # ------------------------------------------------------------------
    # Summary
    # ------------------------------------------------------------------
    print("\n" + "=" * 60)
    total = PASSED + FAILED
    print(f"Results: {PASSED}/{total} passed, {FAILED} failed")
    if FAILED == 0:
        print("M4.01 PHASE B1 SMOKE TESTS: ALL PASSED")
    else:
        print("M4.01 PHASE B1 SMOKE TESTS: FAILURES DETECTED")
    print("=" * 60)

    sys.exit(1 if FAILED > 0 else 0)


if __name__ == "__main__":
    main()
