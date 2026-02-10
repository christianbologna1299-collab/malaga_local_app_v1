"""
Sheet 8: Quantitative Tearsheet — Returns, volatility, drawdown, beta, correlations.
"""

from features.equity.excel.styles import (
    apply_title_row, apply_header_row, apply_freeze_panes, set_column_widths,
    BODY_FONT, SUBTOTAL_FONT, HIGHLIGHT_FONT,
    PERCENT_FORMAT, RATIO_FORMAT, PRICE_FORMAT,
    RIGHT_ALIGN, CENTER_ALIGN, THIN_BORDER, BOTTOM_BORDER,
    HEADER_FILL, HEADER_FONT, LIGHT_GRAY,
)
from openpyxl.styles import PatternFill, Font


def build_quant_sheet(ws, snapshot, outputs):
    """Populate quantitative tearsheet with returns and risk metrics."""
    qm = outputs["quant_metrics"]
    ticker = snapshot["ticker"]

    set_column_widths(ws, {"A": 2, "B": 28, "C": 16, "D": 16, "E": 16})

    # Title
    ws.merge_cells("B1:E1")
    ws["B1"] = "Quantitative Tearsheet"
    apply_title_row(ws, 1, 5)
    apply_freeze_panes(ws, "A3")

    # Key Metrics table header
    ws.cell(row=3, column=2, value="Key Metrics").font = HIGHLIGHT_FONT
    ws.cell(row=3, column=2).border = BOTTOM_BORDER

    # Column headers
    r = 4
    for j, hdr in enumerate(["Metric", ticker, "SPY", "ICLN"]):
        c = ws.cell(row=r, column=2 + j, value=hdr)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER_ALIGN

    # Metrics rows
    metrics = [
        ("Total Return (1Y)", qm["ticker_return"], qm["spy_return"], qm["icln_return"], PERCENT_FORMAT),
        ("Annualized Volatility", qm["ticker_volatility"], qm["spy_volatility"], qm["icln_volatility"], PERCENT_FORMAT),
        ("Sharpe Ratio", qm["ticker_sharpe"], qm["spy_sharpe"], qm["icln_sharpe"], RATIO_FORMAT),
        ("Max Drawdown", qm["ticker_max_drawdown"], qm["spy_max_drawdown"], qm["icln_max_drawdown"], PERCENT_FORMAT),
        ("Current Price", snapshot["current_price"], None, None, PRICE_FORMAT),
    ]
    for i, (label, v1, v2, v3, fmt) in enumerate(metrics):
        row = 5 + i
        ws.cell(row=row, column=2, value=label).font = BODY_FONT
        c = ws.cell(row=row, column=3, value=v1)
        c.number_format = fmt
        c.alignment = RIGHT_ALIGN
        if v2 is not None:
            c2 = ws.cell(row=row, column=4, value=v2)
            c2.number_format = fmt
            c2.alignment = RIGHT_ALIGN
        if v3 is not None:
            c3 = ws.cell(row=row, column=5, value=v3)
            c3.number_format = fmt
            c3.alignment = RIGHT_ALIGN

    # Risk Metrics section
    r = 12
    ws.cell(row=r, column=2, value="Risk Metrics").font = HIGHLIGHT_FONT
    ws.cell(row=r, column=2).border = BOTTOM_BORDER

    risk_items = [
        ("Beta (vs SPY)", qm["beta"], RATIO_FORMAT),
        ("Correlation (vs SPY)", qm["correlation_spy"], RATIO_FORMAT),
        ("Correlation (vs ICLN)", qm["correlation_icln"], RATIO_FORMAT),
    ]
    for i, (label, value, fmt) in enumerate(risk_items):
        row = 13 + i
        ws.cell(row=row, column=2, value=label).font = BODY_FONT
        c = ws.cell(row=row, column=3, value=value)
        c.number_format = fmt
        c.alignment = RIGHT_ALIGN

    # Chart placeholder
    r = 17
    ws.cell(row=r, column=2, value="Chart: Cumulative Returns (placeholder)").font = Font(
        name="Calibri", size=11, color=LIGHT_GRAY, italic=True,
    )
