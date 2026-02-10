"""
Sheet 5: Statement of Cash Flows — Operating + Investing + Financing.
"""

from features.equity.excel.styles import (
    apply_title_row, apply_header_row, apply_freeze_panes, set_column_widths,
    BODY_FONT, SUBTOTAL_FONT, GRAND_TOTAL_FONT,
    NEGATIVE_CURRENCY, PERCENT_FORMAT,
    RIGHT_ALIGN, THIN_BORDER, DOUBLE_BOTTOM_BORDER,
)
from openpyxl.utils import get_column_letter


def build_cash_flow_sheet(ws, snapshot, outputs):
    """Populate Cash Flow statement from projected data."""
    proj = outputs["projected_cash_flow"]
    years = proj["years"]
    n_cols = len(years)
    data_start_col = 3

    widths = {"A": 2, "B": 30}
    for i in range(n_cols):
        widths[get_column_letter(data_start_col + i)] = 16
    set_column_widths(ws, widths)

    # Title
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=data_start_col + n_cols - 1)
    ws["B1"] = "Statement of Cash Flows"
    apply_title_row(ws, 1, data_start_col + n_cols - 1)

    # Column headers
    ws.cell(row=3, column=2, value="$ in millions")
    for i, yr in enumerate(years):
        ws.cell(row=3, column=data_start_col + i, value=f"FY{yr}E").alignment = RIGHT_ALIGN
    apply_header_row(ws, 3, data_start_col + n_cols - 1)
    apply_freeze_panes(ws, "C4")

    def _write_row(row, label, vals, bold=False, border=None):
        ws.cell(row=row, column=2, value=label).font = SUBTOTAL_FONT if bold else BODY_FONT
        for i, v in enumerate(vals):
            c = ws.cell(row=row, column=data_start_col + i, value=v)
            c.number_format = NEGATIVE_CURRENCY
            c.alignment = RIGHT_ALIGN
            c.font = SUBTOTAL_FONT if bold else BODY_FONT
            if border:
                c.border = border

    # Net Income
    r = 5
    _write_row(r, "Net Income", proj["net_income"])

    # Operating Activities
    r = 7
    _write_row(r, "Depreciation & Amortization", proj["da"])

    # Compute working capital changes
    pbs = outputs["projected_balance_sheet"]
    bs = snapshot["balance_sheet"]
    delta_ar = []
    delta_ap = []
    delta_def = []
    for i in range(len(years)):
        if i == 0:
            delta_ar.append(pbs["ar"][i] - bs["ar"][-1])
            delta_ap.append(pbs["ap"][i] - bs["ap"][-1])
            delta_def.append(pbs["deferred_rev"][i] - bs["deferred_rev"][-1])
        else:
            delta_ar.append(pbs["ar"][i] - pbs["ar"][i - 1])
            delta_ap.append(pbs["ap"][i] - pbs["ap"][i - 1])
            delta_def.append(pbs["deferred_rev"][i] - pbs["deferred_rev"][i - 1])

    r = 8
    _write_row(r, "Change in Accounts Receivable", [-v for v in delta_ar])
    r = 9
    _write_row(r, "Change in Accounts Payable", delta_ap)
    r = 10
    _write_row(r, "Change in Deferred Revenue", delta_def)
    r = 11
    _write_row(r, "Operating Cash Flow", proj["operating_cf"], bold=True, border=THIN_BORDER)

    # Investing
    r = 13
    _write_row(r, "Capital Expenditures", proj["investing_cf"])
    r = 14
    _write_row(r, "Investing Cash Flow", proj["investing_cf"], bold=True, border=THIN_BORDER)

    # Financing
    r = 16
    _write_row(r, "Financing Cash Flow", proj["financing_cf"], bold=True, border=THIN_BORDER)

    # Net Cash Flow
    r = 18
    _write_row(r, "Net Cash Flow", proj["net_cash_flow"], bold=True, border=DOUBLE_BOTTOM_BORDER)
    ws.cell(row=18, column=2).font = GRAND_TOTAL_FONT
