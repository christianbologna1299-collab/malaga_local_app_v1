#!/usr/bin/env python3
"""
M4.01 Phase B4.2 Smoke Test — Banker-Grade Polish.

Verifies:
  - Tab colors on all 5 sheets
  - Freeze panes set on Overview (A2) and Scenarios (A2)
  - Print setup on Overview and Scenarios (landscape, fitToPage, print area, margins)
  - Header/Footer text (Banker Analytics — Confidential, session_id, timestamp)
  - Conditional formatting on Scenarios delta column (2 rules)
  - Input highlight fill on Scenarios editable cells
  - Cell comments on Loan Inputs (B3:B10) and Scenarios (B4, B5)
  - Save/reload preserves all settings
  - Backward compat: named ranges, formulas, protection intact

Usage: python scripts/smoke_m4_excel_phase_b4_2.py
Exit code 0 = PASS, nonzero = FAIL
"""

import sys
import os
import tempfile
from datetime import date

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

PASSED = 0
FAILED = 0


def check(name, condition, detail=""):
    global PASSED, FAILED
    if condition:
        PASSED += 1
        print(f"  PASS  {name}")
    else:
        FAILED += 1
        print(f"  FAIL  {name} — {detail}")


def main():
    print("=" * 60)
    print("M4.01 Phase B4.2 Smoke Tests — Banker-Grade Polish")
    print("=" * 60)

    # ------------------------------------------------------------------
    # 1. Imports
    # ------------------------------------------------------------------
    print("\n[1] Imports")
    try:
        from features.excel_export import (
            LoanInputs, ScenarioConfig, SeriesData, ExportMeta,
            _build_skeleton_workbook,
            _apply_b42_polish, _apply_print_setup,
            _apply_conditional_formatting, _apply_tab_colors,
            _apply_input_notes,
            _TAB_COLOR_PRIMARY, _TAB_COLOR_ACCENT, _TAB_COLOR_NEUTRAL,
            _CF_RED_FILL, _CF_GREEN_FILL, _INPUT_HIGHLIGHT_FILL,
            _OV_PRINT_AREA, _SC_PRINT_AREA, _SC_DELTA_RANGE,
            _PRINT_MARGIN_LR, _PRINT_MARGIN_TB, _PRINT_MARGIN_HF,
            _SC_INPUT_DATA_START,
        )
        check("B4.2 imports", True)
    except ImportError as e:
        check("B4.2 imports", False, str(e))
        print("\nCannot continue. Aborting.")
        sys.exit(1)

    try:
        import openpyxl
        check("openpyxl importable", True)
    except ImportError as e:
        check("openpyxl importable", False, str(e))
        sys.exit(1)

    check("_apply_b42_polish is callable", callable(_apply_b42_polish))
    check("_apply_print_setup is callable", callable(_apply_print_setup))
    check("_apply_conditional_formatting is callable",
          callable(_apply_conditional_formatting))
    check("_apply_tab_colors is callable", callable(_apply_tab_colors))
    check("_apply_input_notes is callable", callable(_apply_input_notes))

    # ------------------------------------------------------------------
    # 2. Build workbook
    # ------------------------------------------------------------------
    print("\n[2] Build workbook")
    loan = LoanInputs(
        principal=100000, start_date="2023-01-01", term_months=60,
        rate=6.25, payment_frequency="Monthly",
    )
    scenario = ScenarioConfig(rate_shock_bps=0, balance_shock_pct=0.0,
                              stress_toggle=False)
    series = SeriesData(dates=["2023-01-01"], balances=[100000], rates=[6.25])
    meta = ExportMeta(
        session_id="test-sess-b42",
        data_points=12,
        timestamp="2023-06-15 10:30:00",
    )
    wb = _build_skeleton_workbook(loan, scenario, series, meta)
    check("Workbook created", wb is not None)
    check("5 sheets", len(wb.sheetnames) == 5)

    # ------------------------------------------------------------------
    # 3. Tab colors
    # ------------------------------------------------------------------
    print("\n[3] Tab colors")
    expected_colors = {
        "Overview": _TAB_COLOR_PRIMARY,
        "Loan Inputs": _TAB_COLOR_ACCENT,
        "Amortization": _TAB_COLOR_NEUTRAL,
        "Scenarios": _TAB_COLOR_ACCENT,
        "Charts": _TAB_COLOR_NEUTRAL,
    }
    for name, expected in expected_colors.items():
        ws = wb[name]
        tc = ws.sheet_properties.tabColor
        # tabColor may be a Color object or string
        color_str = str(tc.rgb if hasattr(tc, "rgb") else tc) if tc else ""
        check(f"'{name}' tab color set",
              expected.upper() in color_str.upper(),
              f"got {color_str}")

    # ------------------------------------------------------------------
    # 4. Freeze panes
    # ------------------------------------------------------------------
    print("\n[4] Freeze panes")
    check("Overview freeze = A2",
          wb["Overview"].freeze_panes == "A2",
          f"got {wb['Overview'].freeze_panes}")
    check("Loan Inputs freeze = A3",
          wb["Loan Inputs"].freeze_panes == "A3",
          f"got {wb['Loan Inputs'].freeze_panes}")
    check("Amortization freeze = A3",
          wb["Amortization"].freeze_panes == "A3",
          f"got {wb['Amortization'].freeze_panes}")
    check("Scenarios freeze = A2",
          wb["Scenarios"].freeze_panes == "A2",
          f"got {wb['Scenarios'].freeze_panes}")

    # ------------------------------------------------------------------
    # 5. Print setup — Overview
    # ------------------------------------------------------------------
    print("\n[5] Print setup — Overview")
    ws_ov = wb["Overview"]
    check("Overview orientation = landscape",
          ws_ov.page_setup.orientation == "landscape",
          f"got {ws_ov.page_setup.orientation}")
    check("Overview fitToWidth = 1",
          ws_ov.page_setup.fitToWidth == 1,
          f"got {ws_ov.page_setup.fitToWidth}")
    check("Overview fitToHeight = 1",
          ws_ov.page_setup.fitToHeight == 1,
          f"got {ws_ov.page_setup.fitToHeight}")
    check("Overview fitToPage enabled",
          ws_ov.sheet_properties.pageSetUpPr is not None
          and ws_ov.sheet_properties.pageSetUpPr.fitToPage is True,
          f"got {getattr(ws_ov.sheet_properties.pageSetUpPr, 'fitToPage', None)}")
    check("Overview print area set",
          ws_ov.print_area is not None and len(ws_ov.print_area) > 0,
          f"got {ws_ov.print_area}")
    check("Overview left margin = 0.5",
          abs(ws_ov.page_margins.left - _PRINT_MARGIN_LR) < 0.01,
          f"got {ws_ov.page_margins.left}")

    # ------------------------------------------------------------------
    # 6. Print setup — Scenarios
    # ------------------------------------------------------------------
    print("\n[6] Print setup — Scenarios")
    ws_sc = wb["Scenarios"]
    check("Scenarios orientation = landscape",
          ws_sc.page_setup.orientation == "landscape",
          f"got {ws_sc.page_setup.orientation}")
    check("Scenarios fitToPage enabled",
          ws_sc.sheet_properties.pageSetUpPr is not None
          and ws_sc.sheet_properties.pageSetUpPr.fitToPage is True)
    check("Scenarios print area set",
          ws_sc.print_area is not None and len(ws_sc.print_area) > 0,
          f"got {ws_sc.print_area}")

    # ------------------------------------------------------------------
    # 7. Headers / Footers
    # ------------------------------------------------------------------
    print("\n[7] Headers / Footers")
    ov_hdr = ws_ov.oddHeader.center.text if ws_ov.oddHeader.center else ""
    check("Overview header has 'Banker Analytics'",
          "Banker Analytics" in (ov_hdr or ""),
          f"got '{ov_hdr}'")
    check("Overview header has 'Confidential'",
          "Confidential" in (ov_hdr or ""),
          f"got '{ov_hdr}'")

    ov_ftr_l = ws_ov.oddFooter.left.text if ws_ov.oddFooter.left else ""
    check("Overview footer has session_id",
          "test-sess-b42" in (ov_ftr_l or ""),
          f"got '{ov_ftr_l}'")

    ov_ftr_r = ws_ov.oddFooter.right.text if ws_ov.oddFooter.right else ""
    check("Overview footer has timestamp",
          "2023-06-15" in (ov_ftr_r or ""),
          f"got '{ov_ftr_r}'")

    sc_hdr = ws_sc.oddHeader.center.text if ws_sc.oddHeader.center else ""
    check("Scenarios header has 'Confidential'",
          "Confidential" in (sc_hdr or ""),
          f"got '{sc_hdr}'")

    # ------------------------------------------------------------------
    # 8. Conditional formatting
    # ------------------------------------------------------------------
    print("\n[8] Conditional formatting")
    cf_rules = list(ws_sc.conditional_formatting)
    total_rules = sum(len(cf.rules) for cf in cf_rules)
    check("Scenarios has >= 2 conditional formatting rules",
          total_rules >= 2,
          f"got {total_rules} rules across {len(cf_rules)} entries")

    # ------------------------------------------------------------------
    # 9. Input highlight
    # ------------------------------------------------------------------
    print("\n[9] Input highlight")
    for addr in [f"B{_SC_INPUT_DATA_START}", f"B{_SC_INPUT_DATA_START + 1}"]:
        cell = ws_sc[addr]
        has_fill = (cell.fill is not None
                    and cell.fill.patternType == "solid"
                    and cell.fill.start_color is not None)
        check(f"Scenarios {addr} has highlight fill", has_fill,
              f"patternType={getattr(cell.fill, 'patternType', None)}")

    # ------------------------------------------------------------------
    # 10. Cell comments — Loan Inputs
    # ------------------------------------------------------------------
    print("\n[10] Cell comments — Loan Inputs")
    ws_inp = wb["Loan Inputs"]
    for addr in ["B3", "B4", "B5", "B6", "B7", "B8", "B9", "B10"]:
        has_comment = ws_inp[addr].comment is not None
        check(f"Loan Inputs {addr} has comment", has_comment)

    # ------------------------------------------------------------------
    # 11. Cell comments — Scenarios
    # ------------------------------------------------------------------
    print("\n[11] Cell comments — Scenarios")
    for addr in [f"B{_SC_INPUT_DATA_START}", f"B{_SC_INPUT_DATA_START + 1}"]:
        has_comment = ws_sc[addr].comment is not None
        check(f"Scenarios {addr} has comment", has_comment)

    # Verify comment content
    sc_b4_comment = ws_sc[f"B{_SC_INPUT_DATA_START}"].comment.text
    check("Scenarios B4 comment mentions 'basis points'",
          "basis points" in sc_b4_comment.lower() or "bps" in sc_b4_comment.lower(),
          f"got '{sc_b4_comment}'")

    # ------------------------------------------------------------------
    # 12. Save and reload
    # ------------------------------------------------------------------
    print("\n[12] Save and reload")
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name
    try:
        wb.save(tmp_path)
        check("Workbook saves", True)

        wb2 = openpyxl.load_workbook(tmp_path)
        check("Workbook reloads", True)
        check("5 sheets after reload", len(wb2.sheetnames) == 5)

        # Tab colors survive
        tc2 = wb2["Overview"].sheet_properties.tabColor
        tc2_str = str(tc2.rgb if hasattr(tc2, "rgb") else tc2) if tc2 else ""
        check("Overview tab color survives",
              _TAB_COLOR_PRIMARY.upper() in tc2_str.upper(),
              f"got {tc2_str}")

        # Freeze panes survive
        check("Overview freeze survives",
              wb2["Overview"].freeze_panes == "A2",
              f"got {wb2['Overview'].freeze_panes}")
        check("Scenarios freeze survives",
              wb2["Scenarios"].freeze_panes == "A2",
              f"got {wb2['Scenarios'].freeze_panes}")

        # Print area survives
        ws2_ov = wb2["Overview"]
        check("Overview print area survives",
              ws2_ov.print_area is not None and len(ws2_ov.print_area) > 0,
              f"got {ws2_ov.print_area}")

        # Header survives
        hdr2 = ws2_ov.oddHeader.center.text if ws2_ov.oddHeader.center else ""
        check("Overview header survives reload",
              "Banker Analytics" in (hdr2 or ""),
              f"got '{hdr2}'")

        # Conditional formatting survives
        ws2_sc = wb2["Scenarios"]
        cf2 = list(ws2_sc.conditional_formatting)
        total_rules2 = sum(len(cf.rules) for cf in cf2)
        check("Scenarios CF rules survive reload",
              total_rules2 >= 2,
              f"got {total_rules2} rules across {len(cf2)} entries")

        # Comments survive
        ws2_inp = wb2["Loan Inputs"]
        check("Loan Inputs B4 comment survives",
              ws2_inp["B4"].comment is not None)
        check("Scenarios B4 comment survives",
              ws2_sc[f"B{_SC_INPUT_DATA_START}"].comment is not None)

        # Protection survives (backward compat)
        for name in ["Overview", "Loan Inputs", "Amortization",
                      "Scenarios", "Charts"]:
            check(f"'{name}' protection survives B4.2",
                  wb2[name].protection.sheet is True)

        wb2.close()
    finally:
        os.unlink(tmp_path)

    # ------------------------------------------------------------------
    # 13. Backward compat
    # ------------------------------------------------------------------
    print("\n[13] Backward compat")
    # Named ranges
    defined_names = list(wb.defined_names.keys())
    check("15 named ranges present", len(defined_names) == 15,
          f"got {len(defined_names)}")

    # Formulas
    ws_amort = wb["Amortization"]
    a3 = str(ws_amort["A3"].value)
    check("Amort A3 formula intact", a3.startswith("=") and "Num_Periods" in a3,
          f"got {a3}")

    ws_inp2 = wb["Loan Inputs"]
    b16 = str(ws_inp2["B16"].value)
    check("Loan Inputs B16 PMT formula intact", "PMT" in b16.upper(),
          f"got {b16}")

    # Charts
    ws_ch = wb["Charts"]
    check("3 charts intact", len(ws_ch._charts) == 3,
          f"got {len(ws_ch._charts)}")

    # ------------------------------------------------------------------
    # Summary
    # ------------------------------------------------------------------
    print("\n" + "=" * 60)
    total = PASSED + FAILED
    print(f"Results: {PASSED}/{total} passed, {FAILED} failed")
    if FAILED == 0:
        print("M4.01 PHASE B4.2 SMOKE TESTS: ALL PASSED")
    else:
        print("M4.01 PHASE B4.2 SMOKE TESTS: FAILURES DETECTED")
    print("=" * 60)

    sys.exit(1 if FAILED > 0 else 0)


if __name__ == "__main__":
    main()
