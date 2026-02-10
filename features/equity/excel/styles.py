"""
Shared Excel styling constants for banker-grade equity research workbooks.
Consistent fonts, fills, borders, number formats across all 9 sheets.
"""

from openpyxl.styles import (
    PatternFill, Font, Border, Side, Alignment, numbers,
)

# ── Color constants ──
NAVY = "1B2A4A"
DARK_GRAY = "333333"
MEDIUM_GRAY = "808080"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
INPUT_BLUE = "DCE6F1"
GREEN_GOOD = "C6EFCE"
RED_BAD = "FFC7CE"
GOLD_HIGHLIGHT = "FFF2CC"
GREEN_TINT = "E2EFDA"
RED_TINT = "FCE4EC"

# ── Fills ──
HEADER_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color=DARK_GRAY, end_color=DARK_GRAY, fill_type="solid")
ZEBRA_FILL_1 = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
ZEBRA_FILL_2 = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")
INPUT_FILL = PatternFill(start_color=INPUT_BLUE, end_color=INPUT_BLUE, fill_type="solid")
ASSUMPTIONS_FILL = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")
GREEN_FILL = PatternFill(start_color=GREEN_GOOD, end_color=GREEN_GOOD, fill_type="solid")
RED_FILL = PatternFill(start_color=RED_BAD, end_color=RED_BAD, fill_type="solid")
GOLD_FILL = PatternFill(start_color=GOLD_HIGHLIGHT, end_color=GOLD_HIGHLIGHT, fill_type="solid")
GREEN_TINT_FILL = PatternFill(start_color=GREEN_TINT, end_color=GREEN_TINT, fill_type="solid")
RED_TINT_FILL = PatternFill(start_color=RED_TINT, end_color=RED_TINT, fill_type="solid")

# ── Fonts ──
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color=WHITE)
SUBHEADER_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
BODY_FONT = Font(name="Calibri", size=11, color=DARK_GRAY)
SUBTOTAL_FONT = Font(name="Calibri", size=11, bold=True, color=DARK_GRAY)
GRAND_TOTAL_FONT = Font(name="Calibri", size=11, bold=True, color=NAVY)
HIGHLIGHT_FONT = Font(name="Calibri", size=12, bold=True, color=NAVY)
LINK_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")

# ── Borders ──
THIN_SIDE = Side(style="thin", color=DARK_GRAY)
THIN_BORDER = Border(bottom=THIN_SIDE)
BOTTOM_BORDER = Border(bottom=Side(style="medium", color=DARK_GRAY))
DOUBLE_BOTTOM_BORDER = Border(bottom=Side(style="double", color=DARK_GRAY))
INPUT_BORDER = Border(
    left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE,
)

# ── Alignments ──
LEFT_ALIGN = Alignment(horizontal="left", vertical="center")
RIGHT_ALIGN = Alignment(horizontal="right", vertical="center")
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")

# ── Number Formats ──
CURRENCY_FORMAT = '#,##0'
NEGATIVE_CURRENCY = '#,##0;(#,##0)'
PERCENT_FORMAT = '0.00%'
RATIO_FORMAT = '0.00'
DATE_FORMAT = 'YYYY-MM-DD'
PRICE_FORMAT = '#,##0.00'


def apply_header_row(ws, row, max_col):
    """Apply dark header band styling to a row."""
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN


def apply_title_row(ws, row, max_col):
    """Apply title styling (larger font) to a row."""
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = TITLE_FONT
        cell.alignment = LEFT_ALIGN


def apply_zebra_rows(ws, start_row, end_row, max_col):
    """Apply alternating row shading."""
    for r in range(start_row, end_row + 1):
        fill = ZEBRA_FILL_1 if (r - start_row) % 2 == 0 else ZEBRA_FILL_2
        for col in range(1, max_col + 1):
            ws.cell(row=r, column=col).fill = fill


def apply_freeze_panes(ws, cell_ref):
    """Freeze panes at the given cell reference."""
    ws.freeze_panes = cell_ref


def set_column_widths(ws, widths):
    """Set column widths from a dict like {'A': 30, 'B': 15, ...}."""
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width
