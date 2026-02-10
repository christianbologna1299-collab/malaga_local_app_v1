#!/usr/bin/env python3
"""
M4.01 Phase B4.1 Smoke Test — Sheet Protection + Unlocked Inputs.

Verifies:
  - All 5 sheets are protected (ws.protection.sheet == True)
  - Loan Inputs B3:B10 are unlocked (editable input cells)
  - Scenarios B4:B5 are unlocked (shock parameter inputs)
  - Known formula cells remain locked (default)
  - Named ranges, formulas, and data validations survive protection
  - Save/reload preserves protection and lock states

Usage: python scripts/smoke_m4_excel_phase_b4.py
Exit code 0 = PASS, nonzero = FAIL
"""

import sys
import os
import tempfile
from datetime import date

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

PASSED = 0
FAILED = 0


def check(name, condition, detail=""):
    global PASSED, FAILED
    if condition:
        PASSED += 1
        print(f"  PASS  {name}")
    else:
        FAILED += 1
        print(f"  FAIL  {name} — {detail}")


def _is_unlocked(cell) -> bool:
    """Return True if cell's Protection is explicitly locked=False."""
    return cell.protection.locked is False


def _is_locked(cell) -> bool:
    """Return True if cell is locked (default or explicit)."""
    # openpyxl default: Protection(locked=True)
    prot = cell.protection
    return prot is None or prot.locked is True or prot.locked is None


def main():
    print("=" * 60)
    print("M4.01 Phase B4.1 Smoke Tests — Sheet Protection")
    print("=" * 60)

    # ------------------------------------------------------------------
    # 1. Imports
    # ------------------------------------------------------------------
    print("\n[1] Imports")
    try:
        from features.excel_export import (
            LoanInputs, ScenarioConfig, SeriesData, ExportMeta,
            _build_skeleton_workbook, _apply_sheet_protection,
            _unlock_cells, _LOAN_INPUT_CELLS, _SCENARIO_INPUT_CELLS,
            _SC_INPUT_DATA_START,
        )
        check("B4.1 imports", True)
    except ImportError as e:
        check("B4.1 imports", False, str(e))
        print("\nCannot continue. Aborting.")
        sys.exit(1)

    try:
        import openpyxl
        check("openpyxl importable", True)
    except ImportError as e:
        check("openpyxl importable", False, str(e))
        sys.exit(1)

    check("_apply_sheet_protection is callable", callable(_apply_sheet_protection))
    check("_unlock_cells is callable", callable(_unlock_cells))
    check("_LOAN_INPUT_CELLS = B3:B10",
          _LOAN_INPUT_CELLS == [f"B{r}" for r in range(3, 11)],
          f"got {_LOAN_INPUT_CELLS}")
    check("_SCENARIO_INPUT_CELLS = B4,B5",
          _SCENARIO_INPUT_CELLS == [f"B{_SC_INPUT_DATA_START}",
                                    f"B{_SC_INPUT_DATA_START + 1}"],
          f"got {_SCENARIO_INPUT_CELLS}")

    # ------------------------------------------------------------------
    # 2. Build workbook
    # ------------------------------------------------------------------
    print("\n[2] Build workbook")
    loan = LoanInputs(
        principal=100000, start_date="2023-01-01", term_months=60,
        rate=6.25, payment_frequency="Monthly",
    )
    scenario = ScenarioConfig(rate_shock_bps=0, balance_shock_pct=0.0, stress_toggle=False)
    series = SeriesData(dates=["2023-01-01"], balances=[100000], rates=[6.25])
    wb = _build_skeleton_workbook(loan, scenario, series)
    check("Workbook created", wb is not None)
    check("5 sheets", len(wb.sheetnames) == 5)

    # ------------------------------------------------------------------
    # 3. All sheets protected
    # ------------------------------------------------------------------
    print("\n[3] Sheet protection")
    expected_sheets = ["Overview", "Loan Inputs", "Amortization", "Scenarios", "Charts"]
    for name in expected_sheets:
        ws = wb[name]
        check(f"'{name}' is protected", ws.protection.sheet is True,
              f"sheet={ws.protection.sheet}")

    # ------------------------------------------------------------------
    # 4. Loan Inputs — unlocked input cells
    # ------------------------------------------------------------------
    print("\n[4] Loan Inputs — unlocked cells")
    ws_inp = wb["Loan Inputs"]
    for addr in _LOAN_INPUT_CELLS:
        check(f"Loan Inputs {addr} is unlocked", _is_unlocked(ws_inp[addr]),
              f"locked={ws_inp[addr].protection.locked}")

    # ------------------------------------------------------------------
    # 5. Loan Inputs — locked formula/label cells
    # ------------------------------------------------------------------
    print("\n[5] Loan Inputs — locked cells")
    # Header row
    check("Loan Inputs A1 (header) is locked", _is_locked(ws_inp["A1"]))
    # Labels column
    check("Loan Inputs A4 (label) is locked", _is_locked(ws_inp["A4"]))
    # Derived values (formulas, not editable)
    check("Loan Inputs B13 (Periods/Year formula) is locked",
          _is_locked(ws_inp["B13"]),
          f"locked={ws_inp['B13'].protection.locked}")
    check("Loan Inputs B14 (Periodic Rate formula) is locked",
          _is_locked(ws_inp["B14"]),
          f"locked={ws_inp['B14'].protection.locked}")
    check("Loan Inputs B16 (Payment formula) is locked",
          _is_locked(ws_inp["B16"]),
          f"locked={ws_inp['B16'].protection.locked}")

    # ------------------------------------------------------------------
    # 6. Scenarios — unlocked input cells
    # ------------------------------------------------------------------
    print("\n[6] Scenarios — unlocked cells")
    ws_sc = wb["Scenarios"]
    for addr in _SCENARIO_INPUT_CELLS:
        check(f"Scenarios {addr} is unlocked", _is_unlocked(ws_sc[addr]),
              f"locked={ws_sc[addr].protection.locked}")

    # ------------------------------------------------------------------
    # 7. Scenarios — locked formula cells
    # ------------------------------------------------------------------
    print("\n[7] Scenarios — locked cells")
    check("Scenarios A1 (header) is locked", _is_locked(ws_sc["A1"]))
    # Comparison grid formula
    check("Scenarios B8 (Base Principal formula) is locked",
          _is_locked(ws_sc["B8"]),
          f"locked={ws_sc['B8'].protection.locked}")
    check("Scenarios C8 (Shocked Principal formula) is locked",
          _is_locked(ws_sc["C8"]),
          f"locked={ws_sc['C8'].protection.locked}")

    # ------------------------------------------------------------------
    # 8. Amortization — all locked
    # ------------------------------------------------------------------
    print("\n[8] Amortization — locked")
    ws_amort = wb["Amortization"]
    check("Amort A3 (Period formula) is locked", _is_locked(ws_amort["A3"]))
    check("Amort F3 (Balance formula) is locked", _is_locked(ws_amort["F3"]))
    check("Amort C100 (Payment formula) is locked", _is_locked(ws_amort["C100"]))

    # ------------------------------------------------------------------
    # 9. Overview — all locked
    # ------------------------------------------------------------------
    print("\n[9] Overview — locked")
    ws_ov = wb["Overview"]
    check("Overview A1 (header) is locked", _is_locked(ws_ov["A1"]))
    check("Overview B4 (Loan_Principal ref) is locked", _is_locked(ws_ov["B4"]))
    check("Overview B12 (KPI formula) is locked", _is_locked(ws_ov["B12"]))

    # ------------------------------------------------------------------
    # 10. Charts — all locked
    # ------------------------------------------------------------------
    print("\n[10] Charts — locked")
    ws_ch = wb["Charts"]
    check("Charts A1 (header) is locked", _is_locked(ws_ch["A1"]))
    check("Charts A2 (label) is locked", _is_locked(ws_ch["A2"]))

    # ------------------------------------------------------------------
    # 11. Named ranges survive
    # ------------------------------------------------------------------
    print("\n[11] Named ranges survive")
    defined_names = list(wb.defined_names.keys())
    check("15 named ranges present", len(defined_names) == 15,
          f"got {len(defined_names)}")

    # ------------------------------------------------------------------
    # 12. Data validations survive
    # ------------------------------------------------------------------
    print("\n[12] Data validations survive")
    inp_dvs = ws_inp.data_validations.dataValidation
    check("Loan Inputs has data validation(s)", len(inp_dvs) >= 1,
          f"got {len(inp_dvs)}")
    sc_dvs = ws_sc.data_validations.dataValidation
    check("Scenarios has data validation(s)", len(sc_dvs) >= 1,
          f"got {len(sc_dvs)}")

    # ------------------------------------------------------------------
    # 13. Formulas survive
    # ------------------------------------------------------------------
    print("\n[13] Formulas survive")
    a3_val = str(ws_amort["A3"].value)
    check("Amort A3 formula intact", a3_val.startswith("=") and "Num_Periods" in a3_val,
          f"got {a3_val}")
    b16_val = str(ws_inp["B16"].value)
    check("Loan Inputs B16 PMT formula intact", "PMT" in b16_val.upper(),
          f"got {b16_val}")

    # ------------------------------------------------------------------
    # 14. Save and reload
    # ------------------------------------------------------------------
    print("\n[14] Save and reload")
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name
    try:
        wb.save(tmp_path)
        check("Workbook saves", True)

        wb2 = openpyxl.load_workbook(tmp_path)
        check("Workbook reloads", True)
        check("5 sheets after reload", len(wb2.sheetnames) == 5)

        # Protection survives
        for name in expected_sheets:
            ws2 = wb2[name]
            check(f"'{name}' protection survives reload",
                  ws2.protection.sheet is True,
                  f"sheet={ws2.protection.sheet}")

        # Unlocked cells survive
        ws2_inp = wb2["Loan Inputs"]
        check("Loan Inputs B4 unlocked survives reload",
              _is_unlocked(ws2_inp["B4"]),
              f"locked={ws2_inp['B4'].protection.locked}")
        check("Loan Inputs B8 unlocked survives reload",
              _is_unlocked(ws2_inp["B8"]),
              f"locked={ws2_inp['B8'].protection.locked}")

        ws2_sc = wb2["Scenarios"]
        check(f"Scenarios B{_SC_INPUT_DATA_START} unlocked survives reload",
              _is_unlocked(ws2_sc[f"B{_SC_INPUT_DATA_START}"]),
              f"locked={ws2_sc[f'B{_SC_INPUT_DATA_START}'].protection.locked}")

        # Locked cells survive
        check("Loan Inputs B13 locked survives reload",
              _is_locked(ws2_inp["B13"]),
              f"locked={ws2_inp['B13'].protection.locked}")

        ws2_amort = wb2["Amortization"]
        check("Amort A3 locked survives reload",
              _is_locked(ws2_amort["A3"]),
              f"locked={ws2_amort['A3'].protection.locked}")

        # Named ranges survive reload
        reload_names = list(wb2.defined_names.keys())
        check("15 named ranges survive reload", len(reload_names) == 15,
              f"got {len(reload_names)}")

        wb2.close()
    finally:
        os.unlink(tmp_path)

    # ------------------------------------------------------------------
    # Summary
    # ------------------------------------------------------------------
    print("\n" + "=" * 60)
    total = PASSED + FAILED
    print(f"Results: {PASSED}/{total} passed, {FAILED} failed")
    if FAILED == 0:
        print("M4.01 PHASE B4.1 SMOKE TESTS: ALL PASSED")
    else:
        print("M4.01 PHASE B4.1 SMOKE TESTS: FAILURES DETECTED")
    print("=" * 60)

    sys.exit(1 if FAILED > 0 else 0)


if __name__ == "__main__":
    main()
