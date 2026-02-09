#!/usr/bin/env python3
"""
M4.00 Phase A Smoke Test — Excel Export Engine.
Updated for B1 compatibility (formula-driven amortization, date as datetime).
Verifies workbook structure, named ranges, data validation, sheets, formatting.
Usage: python scripts/smoke_m4_excel_phase_a.py
Exit code 0 = PASS, nonzero = FAIL
"""

import sys
import os
import tempfile
from datetime import date

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

PASSED = 0
FAILED = 0


def check(name, condition, detail=""):
    global PASSED, FAILED
    if condition:
        PASSED += 1
        print(f"  PASS  {name}")
    else:
        FAILED += 1
        print(f"  FAIL  {name} — {detail}")


def main():
    print("=" * 60)
    print("M4.00 Phase A Smoke Tests — Excel Export Engine")
    print("=" * 60)

    # ------------------------------------------------------------------
    # 1. Imports
    # ------------------------------------------------------------------
    print("\n[1] Imports")
    try:
        from features.excel_export import (
            LoanInputs,
            ScenarioConfig,
            SeriesData,
            _build_skeleton_workbook,
            build_excel_workbook,
        )
        check("excel_export imports", True)
    except ImportError as e:
        check("excel_export imports", False, str(e))
        print("\nCannot continue without excel_export. Aborting.")
        sys.exit(1)

    try:
        import openpyxl
        check("openpyxl importable", True)
    except ImportError as e:
        check("openpyxl importable", False, str(e))
        print("\nInstall openpyxl: pip install openpyxl>=3.1.2")
        sys.exit(1)

    # ------------------------------------------------------------------
    # 2. Dataclass instantiation
    # ------------------------------------------------------------------
    print("\n[2] Dataclasses")
    loan = LoanInputs(
        principal=100000,
        start_date="2023-01-01",
        term_months=60,
        rate=6.25,
        payment_frequency="Monthly",
    )
    check("LoanInputs instantiates", loan.principal == 100000)

    scenario = ScenarioConfig(
        rate_shock_bps=100,
        balance_shock_pct=5.0,
        stress_toggle=False,
    )
    check("ScenarioConfig instantiates", scenario.rate_shock_bps == 100)

    series = SeriesData(
        dates=["2023-01-01", "2023-02-01", "2023-03-01", "2023-04-01",
               "2023-05-01", "2023-06-01", "2023-07-01", "2023-08-01",
               "2023-09-01", "2023-10-01", "2023-11-01", "2023-12-01"],
        balances=[100000, 99500, 99000, 98500, 98000, 97500,
                  97000, 96500, 96000, 95500, 95000, 94500],
        rates=[6.0, 6.0, 6.1, 6.1, 6.2, 6.2,
               6.3, 6.3, 6.25, 6.25, 6.2, 6.2],
    )
    check("SeriesData instantiates", len(series.dates) == 12)

    # ------------------------------------------------------------------
    # 3. Build skeleton workbook
    # ------------------------------------------------------------------
    print("\n[3] Skeleton workbook")
    wb = _build_skeleton_workbook(loan, scenario, series)
    check("Workbook created", wb is not None)

    # Sheet names
    expected_sheets = ["Overview", "Loan Inputs", "Amortization", "Scenarios", "Charts"]
    check("Sheet count is 5", len(wb.sheetnames) == 5,
          f"got {len(wb.sheetnames)}: {wb.sheetnames}")
    for name in expected_sheets:
        check(f"Sheet '{name}' exists", name in wb.sheetnames,
              f"missing from {wb.sheetnames}")

    check("No default 'Sheet'", "Sheet" not in wb.sheetnames)

    # ------------------------------------------------------------------
    # 4. Named ranges (Phase A originals)
    # ------------------------------------------------------------------
    print("\n[4] Named ranges")
    expected_names = [
        "Loan_Principal",
        "Loan_Rate",
        "Loan_Term_Months",
        "Loan_Start_Date",
        "Loan_Frequency",
    ]
    defined_names = list(wb.defined_names.keys())
    for name in expected_names:
        check(f"Named range '{name}' exists", name in defined_names,
              f"defined names: {defined_names}")

    # Verify values via Loan Inputs sheet
    ws_inputs = wb["Loan Inputs"]
    check("Loan_Principal value = 100000", ws_inputs["B4"].value == 100000,
          f"got {ws_inputs['B4'].value}")
    check("Loan_Rate value = 0.0625", ws_inputs["B5"].value == 0.0625,
          f"got {ws_inputs['B5'].value}")
    check("Loan_Term_Months value = 60", ws_inputs["B6"].value == 60,
          f"got {ws_inputs['B6'].value}")
    # B7 is now a datetime.date (B1 change)
    check("Loan_Start_Date is date 2023-01-01",
          ws_inputs["B7"].value == date(2023, 1, 1),
          f"got {ws_inputs['B7'].value} (type={type(ws_inputs['B7'].value).__name__})")
    check("Loan_Frequency value = 'Monthly'",
          ws_inputs["B8"].value == "Monthly",
          f"got {ws_inputs['B8'].value}")

    # ------------------------------------------------------------------
    # 5. Data validation
    # ------------------------------------------------------------------
    print("\n[5] Data validation")
    # Check Loan Inputs validations
    input_dvs = ws_inputs.data_validations.dataValidation
    freq_dv_found = False
    for dv in input_dvs:
        for cell_range in dv.sqref.ranges:
            if "B8" in str(cell_range):
                freq_dv_found = True
                check("Frequency dropdown has list type", dv.type == "list")
                check("Frequency dropdown values contain 'Monthly'",
                      "Monthly" in str(dv.formula1),
                      f"formula1={dv.formula1}")
    check("Frequency validation on B8 found", freq_dv_found)

    # Check Scenarios validations
    ws_scenarios = wb["Scenarios"]
    scenario_dvs = ws_scenarios.data_validations.dataValidation
    check("Scenarios sheet has validations", len(scenario_dvs) >= 2,
          f"got {len(scenario_dvs)}")

    # ------------------------------------------------------------------
    # 6. Formatting checks
    # ------------------------------------------------------------------
    print("\n[6] Formatting")
    # Freeze panes
    check("Loan Inputs freeze panes = A3",
          str(ws_inputs.freeze_panes) == "A3",
          f"got {ws_inputs.freeze_panes}")

    ws_amort = wb["Amortization"]
    check("Amortization freeze panes = A3",
          str(ws_amort.freeze_panes) == "A3",
          f"got {ws_amort.freeze_panes}")

    # Header fill on Loan Inputs row 1
    header_cell = ws_inputs.cell(row=1, column=1)
    check("Header has fill color",
          header_cell.fill.start_color.rgb is not None and header_cell.fill.start_color.rgb != "00000000")

    # Column widths
    check("Loan Inputs col A width >= 28",
          ws_inputs.column_dimensions["A"].width >= 28,
          f"got {ws_inputs.column_dimensions['A'].width}")

    # Currency format on principal
    check("Principal cell has currency format",
          ws_inputs["B4"].number_format == '#,##0.00',
          f"got {ws_inputs['B4'].number_format}")

    # ------------------------------------------------------------------
    # 7. Amortization structure (B1: formula-driven, 6 columns)
    # ------------------------------------------------------------------
    print("\n[7] Amortization structure")
    # Header row (row 2) — B1 adds Period column
    amort_headers = [ws_amort.cell(row=2, column=c).value for c in range(1, 7)]
    check("Amortization headers correct (6 cols)",
          amort_headers == ["Period", "Date", "Payment", "Interest", "Principal", "Balance"],
          f"got {amort_headers}")

    # Row 3 should contain formulas (B1)
    check("A3 has formula",
          str(ws_amort["A3"].value).startswith("="),
          f"got {ws_amort['A3'].value}")
    check("F3 has formula",
          str(ws_amort["F3"].value).startswith("="),
          f"got {ws_amort['F3'].value}")

    # ------------------------------------------------------------------
    # 8. Save and reload
    # ------------------------------------------------------------------
    print("\n[8] Save and reload")
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name
    try:
        wb.save(tmp_path)
        check("Workbook saves without error", True)

        wb2 = openpyxl.load_workbook(tmp_path)
        check("Workbook reloads without error", True)
        check("Reloaded sheet count = 5", len(wb2.sheetnames) == 5)

        # Named ranges survive save/reload
        reload_names = list(wb2.defined_names.keys())
        for name in expected_names:
            check(f"Named range '{name}' survives reload", name in reload_names)

        wb2.close()
    finally:
        os.unlink(tmp_path)

    # ------------------------------------------------------------------
    # 9. build_excel_workbook function signature
    # ------------------------------------------------------------------
    print("\n[9] build_excel_workbook entry point")
    import inspect
    sig = inspect.signature(build_excel_workbook)
    params = list(sig.parameters.keys())
    check("Has 'mode' param", "mode" in params)
    check("Has 'feature' param", "feature" in params)
    check("Has 'session_id' param", "session_id" in params)
    check("Has 'analysis_id' param", "analysis_id" in params)
    check("Has 'user_id' param", "user_id" in params)
    check("Has 'request' param", "request" in params)

    # ------------------------------------------------------------------
    # Summary
    # ------------------------------------------------------------------
    print("\n" + "=" * 60)
    total = PASSED + FAILED
    print(f"Results: {PASSED}/{total} passed, {FAILED} failed")
    if FAILED == 0:
        print("M4 PHASE A SMOKE TESTS: ALL PASSED")
    else:
        print("M4 PHASE A SMOKE TESTS: FAILURES DETECTED")
    print("=" * 60)

    sys.exit(1 if FAILED > 0 else 0)


if __name__ == "__main__":
    main()
