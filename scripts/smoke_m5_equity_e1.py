"""
Smoke test for Milestone 5 E1: Equity Research Foundation.
Uses FastAPI TestClient — no running uvicorn needed.
Must pass offline in < 10 seconds. Zero network calls.
"""

import sys
import os
import time
import tempfile

# Ensure project root is on path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from httpx import Client
from starlette.testclient import TestClient
from openpyxl import load_workbook

from app import app

EXPECTED_SHEETS = [
    "Overview",
    "Assumptions",
    "Income Statement",
    "Balance Sheet",
    "Cash Flow",
    "DCF",
    "Sensitivity",
    "Quant Tearsheet",
    "Strategy",
]


def test_e1():
    start = time.time()
    errors = []

    with TestClient(app) as client:
        # 1. Landing page loads
        r = client.get("/equity")
        if r.status_code != 200:
            errors.append(f"Landing page failed: {r.status_code}")
        else:
            print("  [PASS] GET /equity -> 200")

        # 2. Create snapshot
        r = client.post("/equity/snapshot", data={"ticker": "AAPL"})
        if r.status_code != 200:
            errors.append(f"Create snapshot failed: {r.status_code} {r.text}")
            print(f"  [FAIL] POST /equity/snapshot -> {r.status_code}")
            _report(errors, start)
            return
        snap = r.json()
        if "snapshot_id" not in snap:
            errors.append(f"Missing snapshot_id in response: {snap}")
        snapshot_id = snap.get("snapshot_id", "")
        print(f"  [PASS] POST /equity/snapshot -> snapshot_id={snapshot_id[:8]}...")

        # 3. Run model
        r = client.post(f"/equity/{snapshot_id}/run-model")
        if r.status_code != 200:
            errors.append(f"Run model failed: {r.status_code} {r.text}")
            print(f"  [FAIL] POST /equity/{snapshot_id}/run-model -> {r.status_code}")
            _report(errors, start)
            return
        run = r.json()
        if "model_run_id" not in run:
            errors.append(f"Missing model_run_id in response: {run}")
        model_run_id = run.get("model_run_id", "")
        print(f"  [PASS] POST /equity/{{snapshot_id}}/run-model -> model_run_id={model_run_id[:8]}...")

        # 4. Export Excel
        r = client.post(f"/equity/{model_run_id}/export-xlsx")
        if r.status_code != 200:
            errors.append(f"Export XLSX failed: {r.status_code} {r.text}")
            print(f"  [FAIL] POST /equity/{{model_run_id}}/export-xlsx -> {r.status_code}")
            _report(errors, start)
            return

        content_type = r.headers.get("content-type", "")
        if "application/" not in content_type:
            errors.append(f"Unexpected content-type: {content_type}")
        print(f"  [PASS] POST /equity/{{model_run_id}}/export-xlsx -> {len(r.content)} bytes")

        # Save and validate workbook
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp.write(r.content)
            tmp_path = tmp.name

        try:
            wb = load_workbook(tmp_path)
            for sheet_name in EXPECTED_SHEETS:
                if sheet_name not in wb.sheetnames:
                    errors.append(f"Missing sheet: {sheet_name}")
                else:
                    ws = wb[sheet_name]
                    if ws.max_row < 3:
                        errors.append(f"Sheet '{sheet_name}' appears empty (max_row={ws.max_row})")
            wb.close()
            if not errors:
                print(f"  [PASS] Workbook has all 9 sheets with data")
        finally:
            os.remove(tmp_path)

        # 5. Report draft
        r = client.get(f"/equity/{model_run_id}/report-draft")
        if r.status_code != 200:
            errors.append(f"Report draft failed: {r.status_code} {r.text}")
            print(f"  [FAIL] GET /equity/{{model_run_id}}/report-draft -> {r.status_code}")
        else:
            report = r.json()
            text = report.get("report_text", "")
            if len(text) < 50:
                errors.append(f"Report text too short: {len(text)} chars")
            else:
                print(f"  [PASS] GET /equity/{{model_run_id}}/report-draft -> {len(text)} chars")

    _report(errors, start)


def _report(errors, start):
    elapsed = time.time() - start
    if elapsed > 10:
        errors.append(f"Took {elapsed:.1f}s (limit: 10s)")

    if errors:
        print(f"\nFAILED ({len(errors)} error(s)) in {elapsed:.1f}s:")
        for e in errors:
            print(f"  - {e}")
        sys.exit(1)
    else:
        print(f"\nPASSED: E1 smoke test completed in {elapsed:.1f}s")
        sys.exit(0)


if __name__ == "__main__":
    print("=" * 60)
    print("M5 E1 Smoke Test: Equity Research Foundation")
    print("=" * 60)
    test_e1()
