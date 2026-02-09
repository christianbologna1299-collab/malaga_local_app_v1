#!/usr/bin/env python3
"""
M4.01 Phase B5 Smoke Test — Native Excel Chart Objects.

Verifies:
  - Charts sheet contains >= 3 chart objects
  - Chart 1: LineChart "Balance Over Time" refs Amortization col F
  - Chart 2: LineChart "Payment Over Time" refs Amortization col C
  - Chart 3: Stacked AreaChart "Payment Breakdown (Interest vs Principal)" refs cols D+E
  - All charts have X/Y axis titles
  - Y-axis uses currency number format
  - Data ranges point to Amortization rows 3-362
  - Chart constants defined (anchors, dimensions, data bounds)
  - Save/reload preserves all charts
  - Backward compat: named ranges, protection, formulas intact

Usage: python scripts/smoke_m4_excel_phase_b5.py
Exit code 0 = PASS, nonzero = FAIL
"""

import sys
import os
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

PASSED = 0
FAILED = 0


def check(name, condition, detail=""):
    global PASSED, FAILED
    if condition:
        PASSED += 1
        print(f"  PASS  {name}")
    else:
        FAILED += 1
        print(f"  FAIL  {name} — {detail}")


def _chart_title_text(chart):
    """Extract plain text from an openpyxl chart title."""
    title = chart.title
    if title is None:
        return None
    if isinstance(title, str):
        return title
    try:
        for p in title.tx.rich.paragraphs:
            for r in p.r:
                return r.t
    except (AttributeError, IndexError):
        pass
    return str(title)


def _series_ref(series):
    """Extract the formula reference string from a chart series."""
    try:
        return str(series.val.numRef.f)
    except AttributeError:
        return str(series.val)


def main():
    print("=" * 60)
    print("M4.01 Phase B5 Smoke Tests — Native Excel Chart Objects")
    print("=" * 60)

    # ------------------------------------------------------------------
    # 1. Imports & constants
    # ------------------------------------------------------------------
    print("\n[1] Imports & constants")
    try:
        from features.excel_export import (
            LoanInputs, ScenarioConfig, SeriesData,
            _build_skeleton_workbook,
            _make_balance_chart, _make_payment_chart,
            _make_payment_breakdown_chart, _build_charts_sheet,
            _CHART_WIDTH, _CHART_HEIGHT, _CHART_STYLE,
            _CHART_ANCHOR_1, _CHART_ANCHOR_2, _CHART_ANCHOR_3,
            _CHART_CURRENCY_FMT, _CHART_DATA_START_ROW, _CHART_DATA_END_ROW,
            _MAX_AMORT_ROWS,
        )
        check("B5 imports", True)
    except ImportError as e:
        check("B5 imports", False, str(e))
        print("\nCannot continue. Aborting.")
        sys.exit(1)

    try:
        import openpyxl
        from openpyxl.chart import LineChart, AreaChart
        check("openpyxl chart imports", True)
    except ImportError as e:
        check("openpyxl chart imports", False, str(e))
        sys.exit(1)

    # Constants
    check("_CHART_WIDTH = 26", _CHART_WIDTH == 26, f"got {_CHART_WIDTH}")
    check("_CHART_HEIGHT = 12", _CHART_HEIGHT == 12, f"got {_CHART_HEIGHT}")
    check("_CHART_ANCHOR_1 = 'E2'", _CHART_ANCHOR_1 == "E2")
    check("_CHART_ANCHOR_2 = 'E18'", _CHART_ANCHOR_2 == "E18")
    check("_CHART_ANCHOR_3 = 'E34'", _CHART_ANCHOR_3 == "E34")
    check("_CHART_DATA_START_ROW = 3", _CHART_DATA_START_ROW == 3)
    check("_CHART_DATA_END_ROW = 362",
          _CHART_DATA_END_ROW == 2 + _MAX_AMORT_ROWS,
          f"got {_CHART_DATA_END_ROW}")
    check("_CHART_CURRENCY_FMT defined",
          _CHART_CURRENCY_FMT is not None and "$" in _CHART_CURRENCY_FMT)

    # ------------------------------------------------------------------
    # 2. Build workbook
    # ------------------------------------------------------------------
    print("\n[2] Build workbook")
    loan = LoanInputs(
        principal=100000, start_date="2023-01-01", term_months=60,
        rate=6.25, payment_frequency="Monthly",
    )
    scenario = ScenarioConfig(rate_shock_bps=0, balance_shock_pct=0.0,
                              stress_toggle=False)
    series = SeriesData(dates=["2023-01-01"], balances=[100000], rates=[6.25])
    wb = _build_skeleton_workbook(loan, scenario, series)
    check("Workbook created", wb is not None)
    check("Charts sheet exists", "Charts" in wb.sheetnames)

    ws_charts = wb["Charts"]
    charts = ws_charts._charts
    check("Charts sheet has >= 3 charts", len(charts) >= 3,
          f"got {len(charts)}")

    chart1, chart2, chart3 = charts[0], charts[1], charts[2]

    # ------------------------------------------------------------------
    # 3. Chart 1 — Balance Over Time (LineChart)
    # ------------------------------------------------------------------
    print("\n[3] Chart 1 — Balance Over Time")
    check("Chart 1 is LineChart", isinstance(chart1, LineChart),
          f"got {type(chart1).__name__}")

    t1 = _chart_title_text(chart1)
    check("Chart 1 title = 'Balance Over Time'",
          t1 == "Balance Over Time", f"got {t1}")

    check("Chart 1 has 1 series", len(chart1.series) == 1,
          f"got {len(chart1.series)}")

    ref1 = _series_ref(chart1.series[0])
    check("Chart 1 refs col F (Balance)", "$F$" in ref1, f"got {ref1}")
    check("Chart 1 refs row 362", "$362" in ref1 or "362" in ref1, f"got {ref1}")

    check("Chart 1 x-axis title = 'Period'",
          chart1.x_axis.title is not None and "Period" in str(chart1.x_axis.title))
    check("Chart 1 y-axis title = 'Balance ($)'",
          chart1.y_axis.title is not None and "Balance" in str(chart1.y_axis.title))
    check("Chart 1 y-axis currency fmt",
          chart1.y_axis.numFmt is not None and "$" in str(chart1.y_axis.numFmt),
          f"got {chart1.y_axis.numFmt}")

    check("Chart 1 width", chart1.width == _CHART_WIDTH, f"got {chart1.width}")
    check("Chart 1 height", chart1.height == _CHART_HEIGHT, f"got {chart1.height}")

    # ------------------------------------------------------------------
    # 4. Chart 2 — Payment Over Time (LineChart)
    # ------------------------------------------------------------------
    print("\n[4] Chart 2 — Payment Over Time")
    check("Chart 2 is LineChart", isinstance(chart2, LineChart),
          f"got {type(chart2).__name__}")

    t2 = _chart_title_text(chart2)
    check("Chart 2 title = 'Payment Over Time'",
          t2 == "Payment Over Time", f"got {t2}")

    check("Chart 2 has 1 series", len(chart2.series) == 1,
          f"got {len(chart2.series)}")

    ref2 = _series_ref(chart2.series[0])
    check("Chart 2 refs col C (Payment)", "$C$" in ref2, f"got {ref2}")
    check("Chart 2 refs row 362", "$362" in ref2 or "362" in ref2, f"got {ref2}")

    check("Chart 2 x-axis title = 'Period'",
          chart2.x_axis.title is not None and "Period" in str(chart2.x_axis.title))
    check("Chart 2 y-axis title = 'Payment ($)'",
          chart2.y_axis.title is not None and "Payment" in str(chart2.y_axis.title))
    check("Chart 2 y-axis currency fmt",
          chart2.y_axis.numFmt is not None and "$" in str(chart2.y_axis.numFmt),
          f"got {chart2.y_axis.numFmt}")

    # ------------------------------------------------------------------
    # 5. Chart 3 — Payment Breakdown (Stacked AreaChart)
    # ------------------------------------------------------------------
    print("\n[5] Chart 3 — Payment Breakdown")
    check("Chart 3 is AreaChart", isinstance(chart3, AreaChart),
          f"got {type(chart3).__name__}")

    t3 = _chart_title_text(chart3)
    check("Chart 3 title = 'Payment Breakdown (Interest vs Principal)'",
          t3 == "Payment Breakdown (Interest vs Principal)", f"got {t3}")

    check("Chart 3 has 2 series", len(chart3.series) == 2,
          f"got {len(chart3.series)}")
    check("Chart 3 grouping = 'stacked'", chart3.grouping == "stacked",
          f"got {chart3.grouping}")

    ref3a = _series_ref(chart3.series[0])
    ref3b = _series_ref(chart3.series[1])
    check("Chart 3 series 1 refs col D (Interest)", "$D$" in ref3a, f"got {ref3a}")
    check("Chart 3 series 2 refs col E (Principal)", "$E$" in ref3b, f"got {ref3b}")

    check("Chart 3 x-axis title = 'Period'",
          chart3.x_axis.title is not None and "Period" in str(chart3.x_axis.title))
    check("Chart 3 y-axis title = 'Amount ($)'",
          chart3.y_axis.title is not None and "Amount" in str(chart3.y_axis.title))
    check("Chart 3 y-axis currency fmt",
          chart3.y_axis.numFmt is not None and "$" in str(chart3.y_axis.numFmt),
          f"got {chart3.y_axis.numFmt}")

    # ------------------------------------------------------------------
    # 6. Save and reload
    # ------------------------------------------------------------------
    print("\n[6] Save and reload")
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name
    try:
        wb.save(tmp_path)
        check("Workbook saves", True)

        wb2 = openpyxl.load_workbook(tmp_path)
        check("Workbook reloads", True)

        ws2 = wb2["Charts"]
        charts2 = ws2._charts
        check("Charts survive reload (>= 3)", len(charts2) >= 3,
              f"got {len(charts2)}")

        t1r = _chart_title_text(charts2[0])
        t2r = _chart_title_text(charts2[1])
        t3r = _chart_title_text(charts2[2])
        check("Chart 1 title survives", t1r == "Balance Over Time", f"got {t1r}")
        check("Chart 2 title survives", t2r == "Payment Over Time", f"got {t2r}")
        check("Chart 3 title survives",
              t3r == "Payment Breakdown (Interest vs Principal)", f"got {t3r}")

        wb2.close()
    finally:
        os.unlink(tmp_path)

    # ------------------------------------------------------------------
    # 7. Backward compat
    # ------------------------------------------------------------------
    print("\n[7] Backward compat")
    defined_names = list(wb.defined_names.keys())
    check("15 named ranges present", len(defined_names) == 15,
          f"got {len(defined_names)}")

    ws_amort = wb["Amortization"]
    a3 = str(ws_amort["A3"].value)
    check("Amort A3 formula intact", a3.startswith("=") and "Num_Periods" in a3)

    for name in ["Overview", "Loan Inputs", "Amortization", "Scenarios", "Charts"]:
        check(f"'{name}' protection intact", wb[name].protection.sheet is True)

    # ------------------------------------------------------------------
    # Summary
    # ------------------------------------------------------------------
    print("\n" + "=" * 60)
    total = PASSED + FAILED
    print(f"Results: {PASSED}/{total} passed, {FAILED} failed")
    if FAILED == 0:
        print("M4.01 PHASE B5 SMOKE TESTS: ALL PASSED")
    else:
        print("M4.01 PHASE B5 SMOKE TESTS: FAILURES DETECTED")
    print("=" * 60)

    sys.exit(1 if FAILED > 0 else 0)


if __name__ == "__main__":
    main()
