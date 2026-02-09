"""
M4.01 Phase B3: Excel Export Engine — Real Excel Chart Objects.

Produces a banker-grade interactive .xlsx workbook with:
  - Overview (Executive Summary): formula-driven loan summary, KPIs, static snapshot
  - Loan Inputs: user-editable inputs + derived values (B1)
  - Amortization: 360-row IF-guarded formula grid with zebra striping (B1)
  - Scenarios: interactive Base vs Shocked comparison with formula-driven deltas (B2)
  - Charts: 3 real chart objects linked to Amortization data (B3)

Entry point: build_excel_workbook(...)
"""

import logging
import os
from dataclasses import dataclass, field
from datetime import datetime, date
from pathlib import Path
from typing import Literal

from openpyxl import Workbook
from openpyxl.chart import LineChart, AreaChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection, numbers
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from starlette.requests import Request

logger = logging.getLogger(__name__)

# Lazy imports for app-level helpers (avoid circular at module level)
_app_helpers_loaded = False
_safe_json_dumps = None
_get_client_ip = None
_policy_log_event = None
_policy_guard = None
_db = None


def _ensure_app_helpers():
    """Lazy-load app-level helpers to avoid circular imports."""
    global _app_helpers_loaded, _safe_json_dumps, _get_client_ip
    global _policy_log_event, _policy_guard, _db
    if _app_helpers_loaded:
        return
    try:
        from app import safe_json_dumps, get_client_ip, db
        from core.policy import policy_log_event, policy_guard
        _safe_json_dumps = safe_json_dumps
        _get_client_ip = get_client_ip
        _policy_log_event = policy_log_event
        _policy_guard = policy_guard
        _db = db
    except ImportError:
        # Fallback for smoke tests or standalone usage
        import json

        def _safe_json_dumps(obj):
            try:
                return json.dumps(obj, default=str)
            except Exception:
                return '{"error":"serialization_failed"}'

        def _get_client_ip(request):
            return None

        def _policy_log_event(**kwargs):
            pass

        def _policy_guard(ctx):
            pass

        _db = None
    _app_helpers_loaded = True


# =============================================================================
# Dataclasses for typed payloads
# =============================================================================

@dataclass
class LoanInputs:
    principal: float
    start_date: str  # ISO YYYY-MM-DD
    term_months: int
    rate: float  # APR percent (e.g. 6.25)
    payment_frequency: str  # "Monthly" initially
    fees: float | None = None
    amort_type: str = "Level Payment"
    interest_only_months: int | None = None


@dataclass
class ScenarioConfig:
    rate_shock_bps: int
    balance_shock_pct: float
    stress_toggle: bool


@dataclass
class SeriesData:
    dates: list[str]  # ISO strings
    balances: list[float]
    rates: list[float]


@dataclass
class ExportMeta:
    """Static metadata captured at export time for the Export Snapshot block."""
    session_id: str | None = None
    data_points: int = 0
    source_filename: str | None = None
    timestamp: str | None = None  # auto-filled at build time if None


# =============================================================================
# Formatting constants
# =============================================================================

_HEADER_FILL = PatternFill(start_color="1B2A4A", end_color="1B2A4A", fill_type="solid")
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_LABEL_FONT = Font(name="Calibri", bold=True, size=10)
_BODY_FONT = Font(name="Calibri", size=10)
_THIN_BORDER = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
_CURRENCY_FMT = '#,##0.00'
_PERCENT_FMT = '0.00%'
_BPS_FMT = '#,##0'

# B1: formula-driven amortization additions
_ZEBRA_FILL = PatternFill(start_color="F2F6FA", end_color="F2F6FA", fill_type="solid")
_DATE_FMT = 'YYYY-MM-DD'
_RATE_PRECISE_FMT = '0.000000%'
_MAX_AMORT_ROWS = 360

# B2: executive summary + scenario deltas
_SECTION_FILL = PatternFill(start_color="E8ECF1", end_color="E8ECF1", fill_type="solid")
_SECTION_FONT = Font(name="Calibri", bold=True, size=10, color="1B2A4A")
_DELTA_CURRENCY_FMT = '+#,##0.00;-#,##0.00;0.00'
_DELTA_PERCENT_FMT = '+0.00%;-0.00%;0.00%'

# Layout constants (row numbers)
_OV_LOAN_SUMMARY_ROW = 3
_OV_LOAN_DATA_START = 4
_OV_KPI_SECTION_ROW = 11
_OV_KPI_DATA_START = 12
_OV_SNAPSHOT_SECTION_ROW = 17
_OV_SNAPSHOT_DATA_START = 18

_SC_INPUT_SECTION_ROW = 3
_SC_INPUT_DATA_START = 4
_SC_GRID_HEADER_ROW = 7
_SC_GRID_DATA_START = 8

# B3: chart constants
_CHART_WIDTH = 22   # cm
_CHART_HEIGHT = 12  # cm
_CHART_STYLE = 10   # clean, minimal
_COLOR_BALANCE = "1B2A4A"   # navy (matches header)
_COLOR_INTEREST = "C0392B"  # red
_COLOR_PRINCIPAL = "27AE60" # green
_COLOR_PAYMENT = "2980B9"   # teal

# B4.1: sheet protection — editable input cell addresses
# Loan Inputs sheet: B3=Amort Type, B4=Principal, B5=Rate, B6=Term,
#   B7=Start Date, B8=Frequency, B9=Fees, B10=IO Months
_LOAN_INPUT_CELLS = [f"B{r}" for r in range(3, 11)]
# Scenarios sheet: B4=Rate Shock BPS, B5=Balance Shock Pct
_SCENARIO_INPUT_CELLS = [f"B{_SC_INPUT_DATA_START}",
                         f"B{_SC_INPUT_DATA_START + 1}"]


# =============================================================================
# Shared helpers
# =============================================================================

def _apply_header_row(ws, columns: list[str], row: int = 1):
    """Apply dark header strip to row with column names."""
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=col_idx, value=col_name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _THIN_BORDER


def _apply_section_label(ws, row: int, text: str, merge_end_col: str = "B"):
    """Apply a section subheader with fill across columns A:merge_end_col."""
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = _SECTION_FONT
    cell.fill = _SECTION_FILL
    cell.border = _THIN_BORDER
    ws.merge_cells(f"A{row}:{merge_end_col}{row}")
    # Apply fill to merged area (openpyxl only styles the top-left)
    merge_col_end = ord(merge_end_col) - ord("A") + 1
    for c in range(2, merge_col_end + 1):
        mcell = ws.cell(row=row, column=c)
        mcell.fill = _SECTION_FILL
        mcell.border = _THIN_BORDER


def _label_value_row(ws, row: int, label: str, value, fmt: str | None = None,
                     num_cols: int = 2):
    """Write a label/value row with standard formatting."""
    ws.cell(row=row, column=1, value=label).font = _LABEL_FONT
    val_cell = ws.cell(row=row, column=2, value=value)
    val_cell.font = _BODY_FONT
    if fmt:
        val_cell.number_format = fmt
    for c in range(1, num_cols + 1):
        ws.cell(row=row, column=c).border = _THIN_BORDER


# =============================================================================
# B4.1: Sheet protection helpers
# =============================================================================

def _unlock_cells(ws, cell_addresses: list[str]):
    """Mark specific cells as unlocked (editable when sheet is protected)."""
    unlocked = Protection(locked=False)
    for addr in cell_addresses:
        ws[addr].protection = unlocked


def _apply_sheet_protection(wb: Workbook):
    """
    Protect all sheets and unlock only designated input cells.

    Editable cells:
      - Loan Inputs B3:B10  (user inputs)
      - Scenarios B4:B5     (shock parameters)
    All other cells remain locked (Excel default).
    No password — protection prevents accidental edits, not security.
    """
    _unlock_cells(wb["Loan Inputs"], _LOAN_INPUT_CELLS)
    _unlock_cells(wb["Scenarios"], _SCENARIO_INPUT_CELLS)

    for ws in wb.worksheets:
        ws.protection.sheet = True
        ws.protection.enable()


# =============================================================================
# B2: Overview sheet builder (Executive Summary)
# =============================================================================

def _build_overview_sheet(ws, meta: ExportMeta | None):
    """
    Build the Overview sheet as an Executive Summary with three sections:
    1. Loan Summary — formula refs to named ranges (updates with inputs)
    2. Computed KPIs — formula-driven from amortization data
    3. Export Snapshot — static values captured at export time
    """
    meta = meta or ExportMeta()

    # Title
    _apply_header_row(ws, ["Banker Analytics — Executive Summary"], row=1)
    ws.merge_cells("A1:D1")
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    # ---- Section 1: Loan Summary (formula-driven) ----
    _apply_section_label(ws, _OV_LOAN_SUMMARY_ROW, "LOAN SUMMARY")

    loan_summary_rows = [
        # (row, label, formula, format)
        (_OV_LOAN_DATA_START,     "Principal",      "=Loan_Principal",     _CURRENCY_FMT),
        (_OV_LOAN_DATA_START + 1, "Rate (APR)",     "=Loan_Rate",          _PERCENT_FMT),
        (_OV_LOAN_DATA_START + 2, "Term (Months)",  "=Loan_Term_Months",   _BPS_FMT),
        (_OV_LOAN_DATA_START + 3, "Frequency",      "=Loan_Frequency",     None),
        (_OV_LOAN_DATA_START + 4, "Start Date",     "=Loan_Start_Date",    _DATE_FMT),
        (_OV_LOAN_DATA_START + 5, "Payment",        "=Payment",            _CURRENCY_FMT),
    ]
    for row, label, formula, fmt in loan_summary_rows:
        _label_value_row(ws, row, label, formula, fmt)

    # ---- Section 2: Computed KPIs (formula-driven) ----
    _apply_section_label(ws, _OV_KPI_SECTION_ROW, "COMPUTED KPIs")

    kpi_rows = [
        (_OV_KPI_DATA_START,     "Total Payments",
         "=Payment*Num_Periods", _CURRENCY_FMT),
        (_OV_KPI_DATA_START + 1, "Total Interest",
         "=Payment*Num_Periods-Loan_Principal", _CURRENCY_FMT),
        (_OV_KPI_DATA_START + 2, "Ending Balance",
         "=INDEX('Amortization'!F:F,Num_Periods+2)", _CURRENCY_FMT),
        (_OV_KPI_DATA_START + 3, "Total Principal Paid",
         "=Loan_Principal-B14", _CURRENCY_FMT),
    ]
    for row, label, formula, fmt in kpi_rows:
        _label_value_row(ws, row, label, formula, fmt)

    # ---- Section 3: Export Snapshot (static) ----
    _apply_section_label(ws, _OV_SNAPSHOT_SECTION_ROW, "EXPORT SNAPSHOT")

    ts = meta.timestamp or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    sid = meta.session_id or "N/A"
    src = meta.source_filename or "Session export"

    snapshot_rows = [
        (_OV_SNAPSHOT_DATA_START,     "Report Generated", ts,               None),
        (_OV_SNAPSHOT_DATA_START + 1, "Session ID",       sid,              None),
        (_OV_SNAPSHOT_DATA_START + 2, "Data Points",      meta.data_points, _BPS_FMT),
        (_OV_SNAPSHOT_DATA_START + 3, "Source",           src,              None),
    ]
    for row, label, value, fmt in snapshot_rows:
        _label_value_row(ws, row, label, value, fmt)

    # Column widths
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 24


# =============================================================================
# B2: Scenarios sheet builder (Interactive Comparison)
# =============================================================================

def _build_scenarios_sheet(ws, scenario: ScenarioConfig, wb: Workbook):
    """
    Build the Scenarios sheet with interactive Base vs Shocked comparison.
    Adds 5 named ranges to the workbook.
    """
    # Title
    _apply_header_row(ws, ["Scenario Analysis"], row=1)
    ws.merge_cells("A1:D1")
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    # ---- Section 1: Scenario Inputs ----
    _apply_section_label(ws, _SC_INPUT_SECTION_ROW, "SCENARIO INPUTS")

    # Rate Shock (bps) — default from scenario config
    _label_value_row(ws, _SC_INPUT_DATA_START, "Rate Shock (bps)",
                     scenario.rate_shock_bps, _BPS_FMT)
    # Balance Shock (%) — default from scenario config
    _label_value_row(ws, _SC_INPUT_DATA_START + 1, "Balance Shock (%)",
                     scenario.balance_shock_pct, _BPS_FMT)

    # Expanded dropdowns (B2 upgrade)
    dv_rate = DataValidation(
        type="list",
        formula1='"-200,-100,0,100,200"',
        allow_blank=False,
    )
    dv_rate.prompt = "Select rate shock in basis points"
    dv_rate.promptTitle = "Rate Shock"
    ws.add_data_validation(dv_rate)
    dv_rate.add(f"B{_SC_INPUT_DATA_START}")

    dv_bal = DataValidation(
        type="list",
        formula1='"-10,-5,0,5,10"',
        allow_blank=False,
    )
    dv_bal.prompt = "Select balance shock percentage"
    dv_bal.promptTitle = "Balance Shock"
    ws.add_data_validation(dv_bal)
    dv_bal.add(f"B{_SC_INPUT_DATA_START + 1}")

    # Named ranges for inputs
    input_names = {
        "Rate_Shock_BPS": f"'Scenarios'!$B${_SC_INPUT_DATA_START}",
        "Balance_Shock_Pct": f"'Scenarios'!$B${_SC_INPUT_DATA_START + 1}",
    }
    for name, ref in input_names.items():
        wb.defined_names.add(DefinedName(name, attr_text=ref))

    # ---- Section 2: Comparison Grid ----
    # Header row
    grid_headers = ["Metric", "Base Case", "Shocked Case", "Delta"]
    for col_idx, col_name in enumerate(grid_headers, start=1):
        cell = ws.cell(row=_SC_GRID_HEADER_ROW, column=col_idx, value=col_name)
        cell.font = _LABEL_FONT
        cell.fill = _SECTION_FILL
        cell.border = _THIN_BORDER
        cell.alignment = Alignment(horizontal="center")

    # Grid data rows
    # Row layout: (row, metric, base_formula, shocked_formula, delta_formula, base_fmt, shocked_fmt, delta_fmt)
    r = _SC_GRID_DATA_START
    grid_rows = [
        (r, "Principal",
         "=Loan_Principal",
         "=Loan_Principal*(1+Balance_Shock_Pct/100)",
         f"=C{r}-B{r}",
         _CURRENCY_FMT, _CURRENCY_FMT, _DELTA_CURRENCY_FMT),
        (r + 1, "Rate (APR)",
         "=Loan_Rate",
         "=Loan_Rate+Rate_Shock_BPS/10000",
         f"=C{r+1}-B{r+1}",
         _PERCENT_FMT, _PERCENT_FMT, _DELTA_PERCENT_FMT),
        (r + 2, "Payment",
         "=Payment",
         "=-PMT(Shocked_Rate/Periods_Per_Year,Num_Periods,Shocked_Principal)",
         f"=C{r+2}-B{r+2}",
         _CURRENCY_FMT, _CURRENCY_FMT, _DELTA_CURRENCY_FMT),
        (r + 3, "Total Payments",
         "=Payment*Num_Periods",
         "=Shocked_Payment*Num_Periods",
         f"=C{r+3}-B{r+3}",
         _CURRENCY_FMT, _CURRENCY_FMT, _DELTA_CURRENCY_FMT),
        (r + 4, "Total Interest",
         "=Payment*Num_Periods-Loan_Principal",
         "=Shocked_Payment*Num_Periods-Shocked_Principal",
         f"=C{r+4}-B{r+4}",
         _CURRENCY_FMT, _CURRENCY_FMT, _DELTA_CURRENCY_FMT),
    ]

    for row, metric, base_f, shock_f, delta_f, base_fmt, shock_fmt, delta_fmt in grid_rows:
        ws.cell(row=row, column=1, value=metric).font = _LABEL_FONT
        ws.cell(row=row, column=1).border = _THIN_BORDER

        base_cell = ws.cell(row=row, column=2, value=base_f)
        base_cell.font = _BODY_FONT
        base_cell.number_format = base_fmt
        base_cell.border = _THIN_BORDER

        shock_cell = ws.cell(row=row, column=3, value=shock_f)
        shock_cell.font = _BODY_FONT
        shock_cell.number_format = shock_fmt
        shock_cell.border = _THIN_BORDER

        delta_cell = ws.cell(row=row, column=4, value=delta_f)
        delta_cell.font = _BODY_FONT
        delta_cell.number_format = delta_fmt
        delta_cell.border = _THIN_BORDER

    # Named ranges for shocked calculations
    calc_names = {
        "Shocked_Principal": f"'Scenarios'!$C${_SC_GRID_DATA_START}",
        "Shocked_Rate": f"'Scenarios'!$C${_SC_GRID_DATA_START + 1}",
        "Shocked_Payment": f"'Scenarios'!$C${_SC_GRID_DATA_START + 2}",
    }
    for name, ref in calc_names.items():
        wb.defined_names.add(DefinedName(name, attr_text=ref))

    # Column widths
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18


# =============================================================================
# B3: Chart factory functions
# =============================================================================

def _make_balance_chart(ws_amort) -> LineChart:
    """Create Balance Over Time line chart from Amortization data."""
    chart = LineChart()
    chart.title = "Balance Over Time"
    chart.y_axis.title = "Balance ($)"
    chart.x_axis.title = "Period"
    chart.style = _CHART_STYLE
    chart.width = _CHART_WIDTH
    chart.height = _CHART_HEIGHT

    # Data: col F (Balance), rows 2-362 (row 2 = header for auto series title)
    data = Reference(ws_amort, min_col=6, min_row=2, max_row=2 + _MAX_AMORT_ROWS)
    cats = Reference(ws_amort, min_col=1, min_row=3, max_row=2 + _MAX_AMORT_ROWS)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    # Style: smooth navy line, no markers
    s = chart.series[0]
    s.graphicalProperties.line.solidFill = _COLOR_BALANCE
    s.smooth = True

    return chart


def _make_payment_breakdown_chart(ws_amort) -> AreaChart:
    """Create Payment Breakdown stacked area chart (Interest + Principal)."""
    chart = AreaChart()
    chart.title = "Payment Breakdown"
    chart.y_axis.title = "Amount ($)"
    chart.x_axis.title = "Period"
    chart.style = _CHART_STYLE
    chart.width = _CHART_WIDTH
    chart.height = _CHART_HEIGHT
    chart.grouping = "stacked"

    # Data: cols D (Interest) and E (Principal), rows 2-362
    data = Reference(ws_amort, min_col=4, max_col=5, min_row=2,
                     max_row=2 + _MAX_AMORT_ROWS)
    cats = Reference(ws_amort, min_col=1, min_row=3, max_row=2 + _MAX_AMORT_ROWS)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    # Style: Interest=red, Principal=green
    chart.series[0].graphicalProperties.solidFill = _COLOR_INTEREST
    chart.series[1].graphicalProperties.solidFill = _COLOR_PRINCIPAL

    return chart


def _make_payment_chart(ws_amort) -> LineChart:
    """Create Payment Over Time line chart from Amortization data."""
    chart = LineChart()
    chart.title = "Payment Over Time"
    chart.y_axis.title = "Payment ($)"
    chart.x_axis.title = "Period"
    chart.style = _CHART_STYLE
    chart.width = _CHART_WIDTH
    chart.height = _CHART_HEIGHT

    # Data: col C (Payment), rows 2-362
    data = Reference(ws_amort, min_col=3, min_row=2, max_row=2 + _MAX_AMORT_ROWS)
    cats = Reference(ws_amort, min_col=1, min_row=3, max_row=2 + _MAX_AMORT_ROWS)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    # Style: smooth teal line, no markers
    s = chart.series[0]
    s.graphicalProperties.line.solidFill = _COLOR_PAYMENT
    s.smooth = True

    return chart


def _build_charts_sheet(wb: Workbook):
    """
    Populate the Charts sheet with 3 real Excel chart objects
    linked to Amortization data. Charts auto-update when inputs change.
    """
    ws_charts = wb["Charts"]
    ws_amort = wb["Amortization"]

    # Title
    _apply_header_row(ws_charts, ["Charts — Amortization Visualization"], row=1)
    ws_charts.merge_cells("A1:E1")
    ws_charts.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    # Chart labels in column A
    ws_charts["A2"] = "Balance Over Time"
    ws_charts["A2"].font = _LABEL_FONT
    ws_charts["A18"] = "Payment Breakdown"
    ws_charts["A18"].font = _LABEL_FONT
    ws_charts["A34"] = "Payment Over Time"
    ws_charts["A34"].font = _LABEL_FONT

    ws_charts.column_dimensions["A"].width = 40

    # Create and place charts
    chart1 = _make_balance_chart(ws_amort)
    ws_charts.add_chart(chart1, "E2")

    chart2 = _make_payment_breakdown_chart(ws_amort)
    ws_charts.add_chart(chart2, "E18")

    chart3 = _make_payment_chart(ws_amort)
    ws_charts.add_chart(chart3, "E34")


# =============================================================================
# Workbook builder
# =============================================================================

def _build_skeleton_workbook(
    loan: LoanInputs,
    scenario: ScenarioConfig,
    series: SeriesData,
    meta: ExportMeta | None = None,
) -> Workbook:
    """
    Build a banker-grade interactive workbook with 5 sheets,
    named ranges, data validation, formula-driven amortization,
    executive summary, scenario comparisons, and professional formatting.

    Args:
        loan: Loan input parameters
        scenario: Scenario configuration
        series: Time series data (used for Overview data-point count)
        meta: Export metadata for the static snapshot block

    Returns:
        openpyxl.Workbook ready to save
    """
    wb = Workbook()

    # ------------------------------------------------------------------
    # Remove default sheet, create named sheets in order
    # ------------------------------------------------------------------
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    ws_overview = wb.create_sheet("Overview", 0)
    ws_inputs = wb.create_sheet("Loan Inputs", 1)
    ws_amort = wb.create_sheet("Amortization", 2)
    ws_scenarios = wb.create_sheet("Scenarios", 3)
    ws_charts = wb.create_sheet("Charts", 4)

    # ==================================================================
    # 1. Loan Inputs sheet (unchanged from B1)
    # ==================================================================
    _apply_header_row(ws_inputs, ["Loan Inputs", "Value", "Notes"], row=1)

    ws_inputs.cell(row=2, column=1, value="Parameter").font = _LABEL_FONT
    ws_inputs.cell(row=2, column=2, value="Current").font = _LABEL_FONT
    ws_inputs.cell(row=2, column=3, value="Description").font = _LABEL_FONT
    for c in range(1, 4):
        ws_inputs.cell(row=2, column=c).border = _THIN_BORDER

    try:
        start_dt = datetime.strptime(loan.start_date, "%Y-%m-%d").date()
    except (ValueError, TypeError):
        start_dt = date(2023, 1, 1)

    input_rows = [
        ("Amortization Type", loan.amort_type, "Amortization method"),
        ("Principal ($)", loan.principal, "Original loan amount"),
        ("Interest Rate (APR)", loan.rate / 100, "Annual percentage rate"),
        ("Term (Months)", loan.term_months, "Loan duration"),
        ("Start Date", start_dt, "Origination date"),
        ("Payment Frequency", loan.payment_frequency, "Payment schedule"),
        ("Fees ($)", loan.fees if loan.fees is not None else 0, "Origination fees"),
        ("Interest-Only Months", loan.interest_only_months or 0, "IO period length"),
    ]
    for r_idx, (label, value, note) in enumerate(input_rows, start=3):
        ws_inputs.cell(row=r_idx, column=1, value=label).font = _LABEL_FONT
        cell_val = ws_inputs.cell(row=r_idx, column=2, value=value)
        cell_val.font = _BODY_FONT
        ws_inputs.cell(row=r_idx, column=3, value=note).font = _BODY_FONT
        for c in range(1, 4):
            ws_inputs.cell(row=r_idx, column=c).border = _THIN_BORDER

    ws_inputs["B4"].number_format = _CURRENCY_FMT
    ws_inputs["B5"].number_format = _PERCENT_FMT
    ws_inputs["B7"].number_format = _DATE_FMT
    ws_inputs["B9"].number_format = _CURRENCY_FMT

    # Derived Values section (B1)
    ws_inputs.cell(row=12, column=1, value="Derived Values").font = _LABEL_FONT
    ws_inputs.cell(row=12, column=2, value="Calculated").font = _LABEL_FONT
    ws_inputs.cell(row=12, column=3, value="Description").font = _LABEL_FONT
    for c in range(1, 4):
        ws_inputs.cell(row=12, column=c).border = _THIN_BORDER

    derived_rows = [
        (13, "Periods / Year",
         '=IF(Loan_Frequency="Monthly",12,IF(Loan_Frequency="Quarterly",4,12))',
         _BPS_FMT, "From frequency selection"),
        (14, "Periodic Rate",
         '=Loan_Rate/Periods_Per_Year',
         _RATE_PRECISE_FMT, "Per-period interest rate"),
        (15, "Total Periods",
         '=Loan_Term_Months/(12/Periods_Per_Year)',
         _BPS_FMT, "Adjusted for frequency"),
        (16, "Payment",
         '=-PMT(Periodic_Rate,Num_Periods,Loan_Principal)',
         _CURRENCY_FMT, "Level payment amount"),
    ]
    for r_idx, label, formula, fmt, note in derived_rows:
        ws_inputs.cell(row=r_idx, column=1, value=label).font = _LABEL_FONT
        cell_f = ws_inputs.cell(row=r_idx, column=2, value=formula)
        cell_f.font = _BODY_FONT
        cell_f.number_format = fmt
        ws_inputs.cell(row=r_idx, column=3, value=note).font = _BODY_FONT
        for c in range(1, 4):
            ws_inputs.cell(row=r_idx, column=c).border = _THIN_BORDER

    ws_inputs.column_dimensions["A"].width = 28
    ws_inputs.column_dimensions["B"].width = 18
    ws_inputs.column_dimensions["C"].width = 22
    ws_inputs.freeze_panes = "A3"

    dv_freq = DataValidation(
        type="list",
        formula1='"Monthly,Quarterly,Semiannual,Annual"',
        allow_blank=False,
    )
    dv_freq.prompt = "Select payment frequency"
    dv_freq.promptTitle = "Frequency"
    ws_inputs.add_data_validation(dv_freq)
    dv_freq.add("B8")

    # Named ranges — Phase A originals
    named_ranges = {
        "Loan_Principal": "'Loan Inputs'!$B$4",
        "Loan_Rate": "'Loan Inputs'!$B$5",
        "Loan_Term_Months": "'Loan Inputs'!$B$6",
        "Loan_Start_Date": "'Loan Inputs'!$B$7",
        "Loan_Frequency": "'Loan Inputs'!$B$8",
    }
    for name, ref in named_ranges.items():
        wb.defined_names.add(DefinedName(name, attr_text=ref))

    # Named ranges — B1 derived values
    derived_names = {
        "Periods_Per_Year": "'Loan Inputs'!$B$13",
        "Periodic_Rate": "'Loan Inputs'!$B$14",
        "Num_Periods": "'Loan Inputs'!$B$15",
        "Payment": "'Loan Inputs'!$B$16",
    }
    for name, ref in derived_names.items():
        wb.defined_names.add(DefinedName(name, attr_text=ref))

    # ==================================================================
    # 2. Amortization sheet — formula-driven (B1, unchanged)
    # ==================================================================
    _apply_header_row(ws_amort, ["Amortization Schedule"], row=1)
    ws_amort.merge_cells("A1:F1")
    ws_amort.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    amort_cols = ["Period", "Date", "Payment", "Interest", "Principal", "Balance"]
    for col_idx, col_name in enumerate(amort_cols, start=1):
        cell = ws_amort.cell(row=2, column=col_idx, value=col_name)
        cell.font = _LABEL_FONT
        cell.border = _THIN_BORDER
        cell.alignment = Alignment(horizontal="center")

    for col_idx, width in enumerate([10, 14, 16, 16, 16, 16], start=1):
        ws_amort.column_dimensions[get_column_letter(col_idx)].width = width

    ws_amort.freeze_panes = "A3"

    # Row 3: Period 1 (special formulas)
    ws_amort["A3"] = '=IF(ROW()-2>Num_Periods,"",ROW()-2)'
    ws_amort["B3"] = '=IF(A3="","",Loan_Start_Date)'
    ws_amort["C3"] = '=IF(A3="","",Payment)'
    ws_amort["D3"] = '=IF(A3="","",Loan_Principal*Periodic_Rate)'
    ws_amort["E3"] = '=IF(A3="","",C3-D3)'
    ws_amort["F3"] = '=IF(A3="","",Loan_Principal-E3)'

    for c in range(1, 7):
        cell = ws_amort.cell(row=3, column=c)
        cell.font = _BODY_FONT
        cell.border = _THIN_BORDER
    ws_amort["A3"].number_format = _BPS_FMT
    ws_amort["B3"].number_format = _DATE_FMT
    ws_amort["C3"].number_format = _CURRENCY_FMT
    ws_amort["D3"].number_format = _CURRENCY_FMT
    ws_amort["E3"].number_format = _CURRENCY_FMT
    ws_amort["F3"].number_format = _CURRENCY_FMT

    # Rows 4..362: Periods 2..360 (general formulas)
    for r in range(4, 3 + _MAX_AMORT_ROWS):
        prev = r - 1
        ws_amort.cell(row=r, column=1, value=f'=IF(ROW()-2>Num_Periods,"",ROW()-2)')
        ws_amort.cell(row=r, column=2, value=f'=IF(A{r}="","",EDATE(B{prev},12/Periods_Per_Year))')
        ws_amort.cell(row=r, column=3, value=f'=IF(A{r}="","",Payment)')
        ws_amort.cell(row=r, column=4, value=f'=IF(A{r}="","",F{prev}*Periodic_Rate)')
        ws_amort.cell(row=r, column=5, value=f'=IF(A{r}="","",C{r}-D{r})')
        ws_amort.cell(row=r, column=6, value=f'=IF(A{r}="","",F{prev}-E{r})')

        is_zebra = (r - 2) % 2 == 0
        for c in range(1, 7):
            cell = ws_amort.cell(row=r, column=c)
            cell.font = _BODY_FONT
            cell.border = _THIN_BORDER
            if is_zebra:
                cell.fill = _ZEBRA_FILL
        ws_amort.cell(row=r, column=1).number_format = _BPS_FMT
        ws_amort.cell(row=r, column=2).number_format = _DATE_FMT
        ws_amort.cell(row=r, column=3).number_format = _CURRENCY_FMT
        ws_amort.cell(row=r, column=4).number_format = _CURRENCY_FMT
        ws_amort.cell(row=r, column=5).number_format = _CURRENCY_FMT
        ws_amort.cell(row=r, column=6).number_format = _CURRENCY_FMT

    wb.defined_names.add(DefinedName("Amort_Data", attr_text="'Amortization'!$A$2:$F$362"))

    # ==================================================================
    # 3. Overview sheet — Executive Summary (B2)
    # ==================================================================
    _build_overview_sheet(ws_overview, meta)

    # ==================================================================
    # 4. Scenarios sheet — Interactive Comparison (B2)
    # ==================================================================
    _build_scenarios_sheet(ws_scenarios, scenario, wb)

    # ==================================================================
    # 5. Charts sheet — real chart objects (B3)
    # ==================================================================
    _build_charts_sheet(wb)

    # ==================================================================
    # 6. Sheet protection + unlocked inputs (B4.1)
    # ==================================================================
    _apply_sheet_protection(wb)

    return wb


# =============================================================================
# Main entry point
# =============================================================================

def build_excel_workbook(
    *,
    mode: Literal["session", "persistent"],
    feature: Literal["snapshot", "simulator", "explain"],
    session_id: str | None = None,
    analysis_id: int | None = None,
    user_id: int | None = None,
    request: Request | None = None,
    _loan: LoanInputs | None = None,
    _scenario: ScenarioConfig | None = None,
    _series: SeriesData | None = None,
) -> str:
    """
    Build and save an Excel workbook for export.

    Args:
        mode: "session" or "persistent"
        feature: "snapshot", "simulator", or "explain"
        session_id: Required for session mode
        analysis_id: Required for persistent mode
        user_id: Required for persistent mode
        request: Starlette Request (for IP + audit)
        _loan, _scenario, _series: Injected payloads (session mode)

    Returns:
        Absolute path to the saved .xlsx file

    Raises:
        ValueError: Missing required args
        PolicyViolation: Policy check failure
    """
    _ensure_app_helpers()

    req_id = getattr(request.state, "request_id", None) if request else None
    ip = _get_client_ip(request) if request else None

    _policy_log_event(
        db=_db,
        user_id=user_id,
        event_kind="export_xlsx_requested",
        details_json=_safe_json_dumps({
            "mode": mode,
            "feature": feature,
            "session_id": session_id,
            "analysis_id": analysis_id,
        }),
        session_id=session_id,
        analysis_id=analysis_id,
        request_id=req_id,
        ip_address=ip,
    )

    try:
        if mode == "persistent":
            if not user_id or not analysis_id:
                raise ValueError("Persistent mode requires user_id and analysis_id")
            _policy_guard({
                "route": "/excel_export/build",
                "mode": "persistent",
                "user_id": user_id,
                "analysis_id": analysis_id,
            })
            raise NotImplementedError(
                "Persistent-mode Excel export will be implemented in Phase B"
            )

        if mode == "session":
            if not session_id:
                raise ValueError("Session mode requires session_id")

        # ---- Build payload ----
        if _loan and _scenario and _series:
            loan = _loan
            scenario = _scenario
            series = _series
        else:
            from core.database_session_manager import DatabaseSessionManager
            sm = DatabaseSessionManager(ttl_minutes=30)
            sess = sm.get_session(session_id)
            if not sess:
                raise ValueError(f"Session not found or expired: {session_id}")

            df = sess["df"]
            dates = df["date"].dt.strftime("%Y-%m-%d").tolist()
            balances = df["balance"].tolist()
            rates = df["rate"].tolist()

            series = SeriesData(dates=dates, balances=balances, rates=rates)
            loan = LoanInputs(
                principal=balances[0] if balances else 0,
                start_date=dates[0] if dates else "1900-01-01",
                term_months=60,
                rate=float(df["rate"].mean()) * 100 if len(df) > 0 else 5.0,
                payment_frequency="Monthly",
            )
            scenario = ScenarioConfig(
                rate_shock_bps=0,
                balance_shock_pct=0.0,
                stress_toggle=False,
            )

        # ---- Build export metadata ----
        meta = ExportMeta(
            session_id=session_id,
            data_points=len(series.dates),
        )

        # ---- Build workbook ----
        wb = _build_skeleton_workbook(loan, scenario, series, meta)

        # ---- Save ----
        exports_dir = Path(__file__).resolve().parent.parent / "exports"
        exports_dir.mkdir(exist_ok=True)

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_feature = feature.replace("/", "_")
        if session_id:
            short_id = session_id[:8]
        elif analysis_id:
            short_id = str(analysis_id)
        else:
            short_id = "export"
        filename = f"banker_analytics_{safe_feature}_{short_id}_{ts}.xlsx"
        filepath = exports_dir / filename
        wb.save(str(filepath))

        file_size = filepath.stat().st_size
        logger.info(f"Excel export saved: {filepath} ({file_size} bytes)")

        _policy_log_event(
            db=_db,
            user_id=user_id,
            event_kind="export_xlsx_success",
            details_json=_safe_json_dumps({
                "filename": filename,
                "file_size": file_size,
                "mode": mode,
                "feature": feature,
            }),
            session_id=session_id,
            analysis_id=analysis_id,
            request_id=req_id,
            ip_address=ip,
        )

        return str(filepath)

    except Exception as e:
        _policy_log_event(
            db=_db,
            user_id=user_id,
            event_kind="export_xlsx_failed",
            details_json=_safe_json_dumps({
                "error": str(e),
                "mode": mode,
                "feature": feature,
            }),
            session_id=session_id,
            analysis_id=analysis_id,
            request_id=req_id,
            ip_address=ip,
        )
        raise
