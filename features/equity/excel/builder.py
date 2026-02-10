"""
EquityWorkbookBuilder — orchestrates creation of the 9-sheet equity research workbook.
"""

import logging
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from features.equity.excel.sheets.overview import build_overview_sheet
from features.equity.excel.sheets.assumptions import build_assumptions_sheet
from features.equity.excel.sheets.income_statement import build_income_statement_sheet
from features.equity.excel.sheets.balance_sheet import build_balance_sheet_sheet
from features.equity.excel.sheets.cash_flow import build_cash_flow_sheet
from features.equity.excel.sheets.dcf import build_dcf_sheet
from features.equity.excel.sheets.sensitivity import build_sensitivity_sheet
from features.equity.excel.sheets.quant import build_quant_sheet
from features.equity.excel.sheets.strategy import build_strategy_sheet

logger = logging.getLogger(__name__)

SHEET_ORDER = [
    "Overview",
    "Assumptions",
    "Income Statement",
    "Balance Sheet",
    "Cash Flow",
    "DCF",
    "Sensitivity",
    "Quant Tearsheet",
    "Strategy",
]

SHEET_BUILDERS = [
    build_overview_sheet,
    build_assumptions_sheet,
    build_income_statement_sheet,
    build_balance_sheet_sheet,
    build_cash_flow_sheet,
    build_dcf_sheet,
    build_sensitivity_sheet,
    build_quant_sheet,
    build_strategy_sheet,
]


class EquityWorkbookBuilder:
    """Builds the unified 9-sheet equity research workbook."""

    def __init__(self, exports_dir: str = "exports"):
        self.exports_dir = Path(exports_dir)
        self.exports_dir.mkdir(exist_ok=True)

    def build(self, snapshot: dict, outputs: dict, model_run_id: str = "") -> str:
        """Build workbook and save to disk.

        Args:
            snapshot: Raw snapshot data from provider
            outputs: Model engine outputs
            model_run_id: ID for filename

        Returns:
            Absolute file path to the generated .xlsx file.
        """
        ticker = snapshot["ticker"]
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        # Add generated_at to outputs for the overview sheet
        outputs["generated_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        wb = Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        for name, builder in zip(SHEET_ORDER, SHEET_BUILDERS):
            ws = wb.create_sheet(title=name)
            try:
                builder(ws, snapshot, outputs)
            except Exception as e:
                logger.error(f"Error building sheet '{name}': {e}")
                # Write error indicator so sheet isn't completely empty
                ws["B1"] = f"Error building {name}: {str(e)[:100]}"

        # Add named ranges
        self._add_named_ranges(wb, outputs)

        # Save
        filename = f"equity_research_{ticker}_{model_run_id}_{timestamp}.xlsx"
        filepath = self.exports_dir / filename
        wb.save(str(filepath))
        wb.close()

        logger.info(f"Equity workbook saved: {filepath} ({filepath.stat().st_size} bytes)")
        return str(filepath.resolve())

    def build_to_bytes(self, snapshot: dict, outputs: dict, model_run_id: str = "") -> tuple:
        """Build workbook and return (filepath, bytes).

        Returns:
            Tuple of (filepath_str, file_bytes)
        """
        filepath = self.build(snapshot, outputs, model_run_id)
        with open(filepath, "rb") as f:
            content = f.read()
        return filepath, content

    def _add_named_ranges(self, wb, outputs):
        """Add Excel named ranges for key cells."""
        named_ranges = {
            # Overview
            "Ticker": "'Overview'!$C$4",
            "CurrentPrice": "'Overview'!$C$11",
            "ImpliedPrice": "'Overview'!$C$18",
            "MarketCap": "'Overview'!$C$12",
            # DCF
            "WACC": "'DCF'!$C$35",
            "TerminalGrowth": "'DCF'!$C$11",
            "EnterpriseValue": "'DCF'!$C$15",
            "EquityValue": "'DCF'!$C$19",
            # Sensitivity
            "SensitivityBase": "'Sensitivity'!$E$8",
        }
        for name, ref in named_ranges.items():
            try:
                wb.defined_names.add(DefinedName(name, attr_text=ref))
            except Exception as e:
                logger.warning(f"Could not add named range '{name}': {e}")
