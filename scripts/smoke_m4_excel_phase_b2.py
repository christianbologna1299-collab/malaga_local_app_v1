#!/usr/bin/env python3
"""
M4.01 Phase B2 Smoke Test — Executive Summary + Scenario Deltas.

Verifies:
  - ExportMeta dataclass, new formatting constants
  - 5 new named ranges (Rate_Shock_BPS, Balance_Shock_Pct, Shocked_*)
  - Overview: Loan Summary formulas, KPI formulas, static Export Snapshot
  - Scenarios: input cells, expanded dropdowns, comparison grid, delta formatting
  - Save/reload preserves all 15 named ranges + formulas
  - Backward compat: Loan Inputs + Amortization unchanged

Usage: python scripts/smoke_m4_excel_phase_b2.py
Exit code 0 = PASS, nonzero = FAIL
"""

import sys
import os
import tempfile
from datetime import date

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

PASSED = 0
FAILED = 0


def check(name, condition, detail=""):
    global PASSED, FAILED
    if condition:
        PASSED += 1
        print(f"  PASS  {name}")
    else:
        FAILED += 1
        print(f"  FAIL  {name} — {detail}")


def main():
    print("=" * 60)
    print("M4.01 Phase B2 Smoke Tests — Exec Summary + Scenario Deltas")
    print("=" * 60)

    # ------------------------------------------------------------------
    # 1. Imports & constants
    # ------------------------------------------------------------------
    print("\n[1] Imports & constants")
    try:
        from features.excel_export import (
            LoanInputs, ScenarioConfig, SeriesData, ExportMeta,
            _build_skeleton_workbook, _DELTA_CURRENCY_FMT, _DELTA_PERCENT_FMT,
            _SECTION_FILL, _SC_GRID_DATA_START, _SC_INPUT_DATA_START,
            _OV_LOAN_DATA_START, _OV_KPI_DATA_START, _OV_SNAPSHOT_DATA_START,
        )
        check("B2 imports", True)
    except ImportError as e:
        check("B2 imports", False, str(e))
        print("\nCannot continue. Aborting.")
        sys.exit(1)

    try:
        import openpyxl
        check("openpyxl importable", True)
    except ImportError as e:
        check("openpyxl importable", False, str(e))
        sys.exit(1)

    check("ExportMeta instantiates",
          ExportMeta(session_id="abc", data_points=12).session_id == "abc")
    check("_DELTA_CURRENCY_FMT has + prefix", "+" in _DELTA_CURRENCY_FMT)
    check("_DELTA_PERCENT_FMT has + prefix", "+" in _DELTA_PERCENT_FMT)

    # ------------------------------------------------------------------
    # 2. Build workbook with meta
    # ------------------------------------------------------------------
    print("\n[2] Build workbook")
    loan = LoanInputs(
        principal=100000, start_date="2023-01-01", term_months=60,
        rate=6.25, payment_frequency="Monthly",
    )
    scenario = ScenarioConfig(rate_shock_bps=0, balance_shock_pct=0.0, stress_toggle=False)
    series = SeriesData(dates=["2023-01-01"], balances=[100000], rates=[6.25])
    meta = ExportMeta(
        session_id="test-sess-1234",
        data_points=12,
        source_filename="test_export.csv",
        timestamp="2023-06-15 10:30:00",
    )
    wb = _build_skeleton_workbook(loan, scenario, series, meta)
    check("Workbook created", wb is not None)
    check("5 sheets", len(wb.sheetnames) == 5)

    # ------------------------------------------------------------------
    # 3. Named ranges (5 new B2 + 10 existing)
    # ------------------------------------------------------------------
    print("\n[3] Named ranges (B2)")
    defined_names = list(wb.defined_names.keys())
    b2_names = [
        "Rate_Shock_BPS", "Balance_Shock_Pct",
        "Shocked_Principal", "Shocked_Rate", "Shocked_Payment",
    ]
    for name in b2_names:
        check(f"Named range '{name}' exists", name in defined_names,
              f"defined names: {defined_names}")

    check("Total named ranges = 15", len(defined_names) == 15,
          f"got {len(defined_names)}: {defined_names}")

    # ------------------------------------------------------------------
    # 4. Overview — Loan Summary (formulas)
    # ------------------------------------------------------------------
    print("\n[4] Overview — Loan Summary")
    ws_ov = wb["Overview"]

    # Title
    check("Overview title contains 'Executive Summary'",
          "Executive Summary" in str(ws_ov.cell(row=1, column=1).value))

    # Section label
    check("Row 3 = LOAN SUMMARY",
          ws_ov.cell(row=3, column=1).value == "LOAN SUMMARY",
          f"got {ws_ov.cell(row=3, column=1).value}")

    # Formula refs
    b4 = str(ws_ov[f"B{_OV_LOAN_DATA_START}"].value)
    check("B4 refs Loan_Principal", "Loan_Principal" in b4, f"got {b4}")

    b5 = str(ws_ov[f"B{_OV_LOAN_DATA_START + 1}"].value)
    check("B5 refs Loan_Rate", "Loan_Rate" in b5, f"got {b5}")

    b6 = str(ws_ov[f"B{_OV_LOAN_DATA_START + 2}"].value)
    check("B6 refs Loan_Term_Months", "Loan_Term_Months" in b6, f"got {b6}")

    b7 = str(ws_ov[f"B{_OV_LOAN_DATA_START + 3}"].value)
    check("B7 refs Loan_Frequency", "Loan_Frequency" in b7, f"got {b7}")

    b8 = str(ws_ov[f"B{_OV_LOAN_DATA_START + 4}"].value)
    check("B8 refs Loan_Start_Date", "Loan_Start_Date" in b8, f"got {b8}")

    b9 = str(ws_ov[f"B{_OV_LOAN_DATA_START + 5}"].value)
    check("B9 refs Payment", "Payment" in b9, f"got {b9}")

    # ------------------------------------------------------------------
    # 5. Overview — Computed KPIs (formulas)
    # ------------------------------------------------------------------
    print("\n[5] Overview — Computed KPIs")
    check("Row 11 = COMPUTED KPIs",
          ws_ov.cell(row=11, column=1).value == "COMPUTED KPIs",
          f"got {ws_ov.cell(row=11, column=1).value}")

    kpi_b12 = str(ws_ov[f"B{_OV_KPI_DATA_START}"].value)
    check("B12 (Total Payments) refs Payment and Num_Periods",
          "Payment" in kpi_b12 and "Num_Periods" in kpi_b12, f"got {kpi_b12}")

    kpi_b13 = str(ws_ov[f"B{_OV_KPI_DATA_START + 1}"].value)
    check("B13 (Total Interest) refs Loan_Principal",
          "Loan_Principal" in kpi_b13, f"got {kpi_b13}")

    kpi_b14 = str(ws_ov[f"B{_OV_KPI_DATA_START + 2}"].value)
    check("B14 (Ending Balance) uses INDEX into Amortization",
          "INDEX" in kpi_b14.upper() and "Amortization" in kpi_b14,
          f"got {kpi_b14}")

    kpi_b15 = str(ws_ov[f"B{_OV_KPI_DATA_START + 3}"].value)
    check("B15 (Total Principal Paid) refs Loan_Principal and B14",
          "Loan_Principal" in kpi_b15 and "B14" in kpi_b15, f"got {kpi_b15}")

    # ------------------------------------------------------------------
    # 6. Overview — Export Snapshot (static)
    # ------------------------------------------------------------------
    print("\n[6] Overview — Export Snapshot")
    check("Row 17 = EXPORT SNAPSHOT",
          ws_ov.cell(row=17, column=1).value == "EXPORT SNAPSHOT",
          f"got {ws_ov.cell(row=17, column=1).value}")

    snap_ts = ws_ov[f"B{_OV_SNAPSHOT_DATA_START}"].value
    check("B18 (timestamp) is static string",
          isinstance(snap_ts, str) and not str(snap_ts).startswith("="),
          f"got {snap_ts} (type={type(snap_ts).__name__})")
    check("B18 matches meta.timestamp", snap_ts == "2023-06-15 10:30:00",
          f"got {snap_ts}")

    snap_sid = ws_ov[f"B{_OV_SNAPSHOT_DATA_START + 1}"].value
    check("B19 (session ID) is static", isinstance(snap_sid, str))
    check("B19 = 'test-sess-1234'", snap_sid == "test-sess-1234",
          f"got {snap_sid}")

    snap_dp = ws_ov[f"B{_OV_SNAPSHOT_DATA_START + 2}"].value
    check("B20 (data points) is integer",
          isinstance(snap_dp, int) and snap_dp == 12,
          f"got {snap_dp} (type={type(snap_dp).__name__})")

    snap_src = ws_ov[f"B{_OV_SNAPSHOT_DATA_START + 3}"].value
    check("B21 (source) = 'test_export.csv'", snap_src == "test_export.csv",
          f"got {snap_src}")

    # ------------------------------------------------------------------
    # 7. Scenarios — Inputs
    # ------------------------------------------------------------------
    print("\n[7] Scenarios — Inputs")
    ws_sc = wb["Scenarios"]

    check("Scenarios title contains 'Scenario Analysis'",
          "Scenario Analysis" in str(ws_sc.cell(row=1, column=1).value))

    check("Row 3 = SCENARIO INPUTS",
          ws_sc.cell(row=3, column=1).value == "SCENARIO INPUTS",
          f"got {ws_sc.cell(row=3, column=1).value}")

    rate_shock_val = ws_sc[f"B{_SC_INPUT_DATA_START}"].value
    check("B4 (Rate Shock) default = 0", rate_shock_val == 0,
          f"got {rate_shock_val}")

    bal_shock_val = ws_sc[f"B{_SC_INPUT_DATA_START + 1}"].value
    check("B5 (Balance Shock) default = 0.0", bal_shock_val == 0.0,
          f"got {bal_shock_val}")

    # Dropdown validations
    sc_dvs = ws_sc.data_validations.dataValidation
    check("Scenarios has 2 validations", len(sc_dvs) == 2,
          f"got {len(sc_dvs)}")

    rate_dv_ok = False
    bal_dv_ok = False
    for dv in sc_dvs:
        formula_str = str(dv.formula1)
        for cell_range in dv.sqref.ranges:
            cell_str = str(cell_range)
            if f"B{_SC_INPUT_DATA_START}" == cell_str:
                rate_dv_ok = "-200" in formula_str and "200" in formula_str
            if f"B{_SC_INPUT_DATA_START + 1}" == cell_str:
                bal_dv_ok = "-10" in formula_str and "10" in formula_str
    check("Rate Shock dropdown has -200..200", rate_dv_ok)
    check("Balance Shock dropdown has -10..10", bal_dv_ok)

    # ------------------------------------------------------------------
    # 8. Scenarios — Comparison Grid
    # ------------------------------------------------------------------
    print("\n[8] Scenarios — Comparison Grid")

    # Header row
    grid_h = [ws_sc.cell(row=7, column=c).value for c in range(1, 5)]
    check("Grid headers correct",
          grid_h == ["Metric", "Base Case", "Shocked Case", "Delta"],
          f"got {grid_h}")

    r = _SC_GRID_DATA_START

    # Row 8: Principal
    b_r8 = str(ws_sc.cell(row=r, column=2).value)
    check("B8 (Base Principal) refs Loan_Principal", "Loan_Principal" in b_r8,
          f"got {b_r8}")

    c_r8 = str(ws_sc.cell(row=r, column=3).value)
    check("C8 (Shocked Principal) refs Balance_Shock_Pct",
          "Balance_Shock_Pct" in c_r8, f"got {c_r8}")

    d_r8 = str(ws_sc.cell(row=r, column=4).value)
    check("D8 (Delta Principal) is formula", d_r8.startswith("="), f"got {d_r8}")

    # Row 9: Rate
    c_r9 = str(ws_sc.cell(row=r + 1, column=3).value)
    check("C9 (Shocked Rate) refs Rate_Shock_BPS",
          "Rate_Shock_BPS" in c_r9, f"got {c_r9}")

    # Row 10: Payment
    b_r10 = str(ws_sc.cell(row=r + 2, column=2).value)
    check("B10 (Base Payment) refs Payment", "Payment" in b_r10, f"got {b_r10}")

    c_r10 = str(ws_sc.cell(row=r + 2, column=3).value)
    check("C10 (Shocked Payment) uses PMT", "PMT" in c_r10.upper(), f"got {c_r10}")
    check("C10 refs Shocked_Rate", "Shocked_Rate" in c_r10, f"got {c_r10}")
    check("C10 refs Shocked_Principal", "Shocked_Principal" in c_r10, f"got {c_r10}")

    d_r10 = str(ws_sc.cell(row=r + 2, column=4).value)
    check("D10 (Delta Payment) is formula", d_r10.startswith("="), f"got {d_r10}")

    # Row 12: Total Interest
    c_r12 = str(ws_sc.cell(row=r + 4, column=3).value)
    check("C12 (Shocked Total Interest) refs Shocked_Payment",
          "Shocked_Payment" in c_r12, f"got {c_r12}")

    d_r12 = str(ws_sc.cell(row=r + 4, column=4).value)
    check("D12 (Delta Interest) is formula", d_r12.startswith("="), f"got {d_r12}")

    # ------------------------------------------------------------------
    # 9. Delta formatting
    # ------------------------------------------------------------------
    print("\n[9] Delta formatting")
    d10_fmt = ws_sc.cell(row=r + 2, column=4).number_format
    check("D10 has +/- currency format", "+" in d10_fmt,
          f"got {d10_fmt}")

    d9_fmt = ws_sc.cell(row=r + 1, column=4).number_format
    check("D9 has +/- percent format", "+" in d9_fmt,
          f"got {d9_fmt}")

    # ------------------------------------------------------------------
    # 10. Save and reload
    # ------------------------------------------------------------------
    print("\n[10] Save and reload")
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name
    try:
        wb.save(tmp_path)
        check("Workbook saves", True)

        wb2 = openpyxl.load_workbook(tmp_path)
        check("Workbook reloads", True)

        # All 15 named ranges survive
        reload_names = list(wb2.defined_names.keys())
        check("15 named ranges survive reload", len(reload_names) == 15,
              f"got {len(reload_names)}: {reload_names}")
        for name in b2_names:
            check(f"'{name}' survives reload", name in reload_names)

        # Overview formulas survive
        ws2_ov = wb2["Overview"]
        b4_reload = str(ws2_ov[f"B{_OV_LOAN_DATA_START}"].value)
        check("Overview B4 formula survives", "Loan_Principal" in b4_reload,
              f"got {b4_reload}")

        b14_reload = str(ws2_ov[f"B{_OV_KPI_DATA_START + 2}"].value)
        check("Overview B14 (Ending Balance) survives", "INDEX" in b14_reload.upper(),
              f"got {b14_reload}")

        # Snapshot static values survive
        snap_ts_reload = ws2_ov[f"B{_OV_SNAPSHOT_DATA_START}"].value
        check("Snapshot timestamp survives as static",
              snap_ts_reload == "2023-06-15 10:30:00",
              f"got {snap_ts_reload}")

        # Scenarios formulas survive
        ws2_sc = wb2["Scenarios"]
        c10_reload = str(ws2_sc.cell(row=r + 2, column=3).value)
        check("Scenarios C10 formula survives", "PMT" in c10_reload.upper(),
              f"got {c10_reload}")

        wb2.close()
    finally:
        os.unlink(tmp_path)

    # ------------------------------------------------------------------
    # 11. Backward compat
    # ------------------------------------------------------------------
    print("\n[11] Backward compat")
    check("5 sheets present", len(wb.sheetnames) == 5)

    # Amortization unchanged
    ws_amort = wb["Amortization"]
    a3 = str(ws_amort["A3"].value)
    check("Amort A3 still formula", a3.startswith("=") and "Num_Periods" in a3,
          f"got {a3}")

    # Loan Inputs unchanged
    ws_inp = wb["Loan Inputs"]
    b13 = str(ws_inp["B13"].value)
    check("Loan Inputs B13 still has Periods_Per_Year formula",
          "Loan_Frequency" in b13, f"got {b13}")

    b7_val = ws_inp["B7"].value
    check("Start date still datetime.date",
          isinstance(b7_val, date), f"got type {type(b7_val).__name__}")

    # ------------------------------------------------------------------
    # 12. ExportMeta defaults
    # ------------------------------------------------------------------
    print("\n[12] ExportMeta defaults")
    wb_no_meta = _build_skeleton_workbook(loan, scenario, series)
    ws_ov2 = wb_no_meta["Overview"]
    snap_sid2 = ws_ov2[f"B{_OV_SNAPSHOT_DATA_START + 1}"].value
    check("No-meta session ID = 'N/A'", snap_sid2 == "N/A", f"got {snap_sid2}")
    snap_src2 = ws_ov2[f"B{_OV_SNAPSHOT_DATA_START + 3}"].value
    check("No-meta source = 'Session export'", snap_src2 == "Session export",
          f"got {snap_src2}")

    # ------------------------------------------------------------------
    # Summary
    # ------------------------------------------------------------------
    print("\n" + "=" * 60)
    total = PASSED + FAILED
    print(f"Results: {PASSED}/{total} passed, {FAILED} failed")
    if FAILED == 0:
        print("M4.01 PHASE B2 SMOKE TESTS: ALL PASSED")
    else:
        print("M4.01 PHASE B2 SMOKE TESTS: FAILURES DETECTED")
    print("=" * 60)

    sys.exit(1 if FAILED > 0 else 0)


if __name__ == "__main__":
    main()
