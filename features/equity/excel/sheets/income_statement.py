"""
Sheet 3: Income Statement — Historic + Projected with assumptions section.
"""

from features.equity.excel.styles import (
    apply_title_row, apply_header_row, apply_freeze_panes, set_column_widths,
    BODY_FONT, SUBTOTAL_FONT, GRAND_TOTAL_FONT, ASSUMPTIONS_FILL,
    CURRENCY_FORMAT, NEGATIVE_CURRENCY, PERCENT_FORMAT,
    LEFT_ALIGN, RIGHT_ALIGN, THIN_BORDER, BOTTOM_BORDER, DOUBLE_BOTTOM_BORDER,
)
from openpyxl.utils import get_column_letter


def build_income_statement_sheet(ws, snapshot, outputs):
    """Populate Income Statement with historic + projected data."""
    hist = snapshot["historic_financials"]
    proj = outputs["projected_income_statement"]
    a = outputs["assumptions"]

    h_years = hist["years"]
    p_years = proj["years"]
    all_years = h_years + p_years
    n_cols = len(all_years)
    data_start_col = 3  # Column C

    # Column widths
    widths = {"A": 2, "B": 28}
    for i in range(n_cols):
        widths[get_column_letter(data_start_col + i)] = 16
    set_column_widths(ws, widths)

    # Row 1: Title
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=data_start_col + n_cols - 1)
    ws["B1"] = "Income Statement"
    apply_title_row(ws, 1, data_start_col + n_cols - 1)

    # Row 3: Column headers
    ws.cell(row=3, column=2, value="$ in millions")
    for i, yr in enumerate(all_years):
        label = f"FY{yr}"
        if yr in p_years:
            label += "E"
        c = ws.cell(row=3, column=data_start_col + i, value=label)
        c.alignment = RIGHT_ALIGN
    apply_header_row(ws, 3, data_start_col + n_cols - 1)
    apply_freeze_panes(ws, "C4")

    def _write_row(row, label, hist_vals, proj_vals, fmt=NEGATIVE_CURRENCY, bold=False):
        ws.cell(row=row, column=2, value=label).font = SUBTOTAL_FONT if bold else BODY_FONT
        for i, v in enumerate(hist_vals + proj_vals):
            c = ws.cell(row=row, column=data_start_col + i, value=v)
            c.number_format = fmt
            c.alignment = RIGHT_ALIGN
            c.font = SUBTOTAL_FONT if bold else BODY_FONT
        if bold:
            for i in range(n_cols + 1):
                ws.cell(row=row, column=2 + i).border = THIN_BORDER

    # Revenue section
    r = 5
    _write_row(r, "Net Revenue", hist["revenue"], proj["revenue"], bold=True)

    r = 7
    _write_row(r, "Cost of Goods Sold", hist["cogs"], proj["cogs"])
    r = 8
    _write_row(r, "Total COGS", hist["cogs"], proj["cogs"], bold=True)

    # Gross Profit
    r = 10
    _write_row(r, "Gross Profit", hist["gross_profit"], proj["gross_profit"], bold=True)
    # Gross margin %
    r = 11
    ws.cell(row=r, column=2, value="Gross Margin %").font = BODY_FONT
    for i, (rev, gp) in enumerate(
        zip(hist["revenue"] + proj["revenue"], hist["gross_profit"] + proj["gross_profit"])
    ):
        c = ws.cell(row=r, column=data_start_col + i, value=gp / rev if rev else 0)
        c.number_format = PERCENT_FORMAT
        c.alignment = RIGHT_ALIGN

    # OpEx
    r = 13
    _write_row(r, "SG&A", hist["sga"], proj["sga"])
    r = 14
    _write_row(r, "Total Operating Expenses", hist["sga"], proj["sga"], bold=True)

    # EBITDA
    r = 16
    _write_row(r, "EBITDA", hist["ebitda"], proj["ebitda"], bold=True)

    # D&A and EBIT
    r = 18
    _write_row(r, "Depreciation & Amortization", hist["da"], proj["da"])
    r = 19
    _write_row(r, "EBIT", hist["ebit"], proj["ebit"], bold=True)

    # Interest, EBT, Tax, Net Income
    r = 21
    _write_row(r, "Interest Expense", hist["interest"], proj["interest"])
    r = 22
    _write_row(r, "Earnings Before Tax", hist["ebt"], proj["ebt"], bold=True)
    r = 23
    _write_row(r, "Income Tax", hist["tax"], proj["tax"])
    r = 24
    _write_row(r, "Net Income", hist["net_income"], proj["net_income"], bold=True)
    # Double underline on Net Income
    for i in range(n_cols):
        ws.cell(row=24, column=data_start_col + i).border = DOUBLE_BOTTOM_BORDER
    ws.cell(row=24, column=2).font = GRAND_TOTAL_FONT

    # Net Income Margin %
    r = 26
    ws.cell(row=r, column=2, value="Net Income Margin %").font = BODY_FONT
    for i, (rev, ni) in enumerate(
        zip(hist["revenue"] + proj["revenue"], hist["net_income"] + proj["net_income"])
    ):
        c = ws.cell(row=r, column=data_start_col + i, value=ni / rev if rev else 0)
        c.number_format = PERCENT_FORMAT
        c.alignment = RIGHT_ALIGN

    # Assumptions section (row 28+)
    r = 28
    ws.cell(row=r, column=2, value="Assumptions").font = SUBTOTAL_FONT
    ws.cell(row=r, column=2).border = BOTTOM_BORDER

    assumptions_items = [
        ("Revenue CAGR", a["revenue_cagr"], PERCENT_FORMAT),
        ("COGS % of Revenue", a["cogs_pct"], PERCENT_FORMAT),
        ("SG&A % of Revenue", a["sga_pct"], PERCENT_FORMAT),
        ("D&A % of Revenue", a["da_pct"], PERCENT_FORMAT),
        ("Tax Rate", a["tax_rate"], PERCENT_FORMAT),
        ("Cost of Debt (Interest Rate)", a["cost_of_debt"], PERCENT_FORMAT),
    ]
    for i, (label, value, fmt) in enumerate(assumptions_items):
        row = 29 + i
        c = ws.cell(row=row, column=2, value=label)
        c.font = BODY_FONT
        c.fill = ASSUMPTIONS_FILL
        vc = ws.cell(row=row, column=3, value=value)
        vc.number_format = fmt
        vc.alignment = RIGHT_ALIGN
        vc.fill = ASSUMPTIONS_FILL
