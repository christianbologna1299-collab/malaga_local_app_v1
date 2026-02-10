"""
Sheet 6: DCF Valuation — UFCF → PV → Enterprise Value → Equity Value → Implied Price + WACC build-up.
"""

from features.equity.excel.styles import (
    apply_title_row, apply_header_row, apply_freeze_panes, set_column_widths,
    BODY_FONT, SUBTOTAL_FONT, GRAND_TOTAL_FONT, HIGHLIGHT_FONT,
    GOLD_FILL,
    CURRENCY_FORMAT, NEGATIVE_CURRENCY, PERCENT_FORMAT, PRICE_FORMAT, RATIO_FORMAT,
    RIGHT_ALIGN, LEFT_ALIGN, THIN_BORDER, BOTTOM_BORDER, DOUBLE_BOTTOM_BORDER,
)
from openpyxl.utils import get_column_letter


def build_dcf_sheet(ws, snapshot, outputs):
    """Populate DCF sheet with full valuation bridge and WACC calculation."""
    dcf = outputs["dcf"]
    wacc_calc = outputs["wacc_calc"]
    proj_years = outputs["projected_income_statement"]["years"]
    n_years = len(proj_years)
    data_start_col = 3

    widths = {"A": 2, "B": 35, "C": 18}
    for i in range(n_years):
        widths[get_column_letter(data_start_col + i)] = 16
    set_column_widths(ws, widths)

    # Title
    last_col = max(data_start_col + n_years - 1, 3)
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=last_col)
    ws["B1"] = "DCF Valuation"
    apply_title_row(ws, 1, last_col)
    apply_freeze_panes(ws, "C4")

    # UFCF row (row 3-4)
    ws.cell(row=3, column=2, value="$ in millions")
    for i, yr in enumerate(proj_years):
        ws.cell(row=3, column=data_start_col + i, value=f"FY{yr}E").alignment = RIGHT_ALIGN
    apply_header_row(ws, 3, last_col)

    ws.cell(row=4, column=2, value="Unlevered Free Cash Flow").font = SUBTOTAL_FONT
    for i, v in enumerate(dcf["ufcf"]):
        c = ws.cell(row=4, column=data_start_col + i, value=v)
        c.number_format = NEGATIVE_CURRENCY
        c.alignment = RIGHT_ALIGN
        c.font = SUBTOTAL_FONT

    # Projection year numbers (row 6)
    ws.cell(row=6, column=2, value="Year").font = BODY_FONT
    for i in range(n_years):
        ws.cell(row=6, column=data_start_col + i, value=i + 1).alignment = RIGHT_ALIGN

    # PV of FCF (row 7)
    ws.cell(row=7, column=2, value="PV of Free Cash Flow").font = BODY_FONT
    for i, v in enumerate(dcf["pv_fcf"]):
        c = ws.cell(row=7, column=data_start_col + i, value=v)
        c.number_format = NEGATIVE_CURRENCY
        c.alignment = RIGHT_ALIGN

    # Valuation Bridge (rows 9-21)
    r = 9
    ws.cell(row=r, column=2, value="Implied Share Price Calculation").font = HIGHLIGHT_FONT
    ws.cell(row=r, column=2).border = BOTTOM_BORDER

    bridge_items = [
        ("Sum of PV of FCF", dcf["sum_pv_fcf"], CURRENCY_FORMAT, True),
        ("Terminal Growth Rate", dcf["terminal_growth"], PERCENT_FORMAT, False),
        ("WACC", dcf["wacc_used"], PERCENT_FORMAT, False),
        ("Terminal Value", dcf["terminal_value"], CURRENCY_FORMAT, False),
        ("PV of Terminal Value", dcf["pv_terminal"], CURRENCY_FORMAT, False),
        ("Enterprise Value", dcf["enterprise_value"], CURRENCY_FORMAT, True),
        ("(+) Cash", dcf["cash"], CURRENCY_FORMAT, False),
        ("(-) Debt", dcf["debt"], CURRENCY_FORMAT, False),
        ("(-) Minority Interest", dcf["minority_interest"], CURRENCY_FORMAT, False),
        ("Equity Value", dcf["equity_value"], CURRENCY_FORMAT, True),
        ("Diluted Shares Outstanding", dcf["shares_outstanding"], CURRENCY_FORMAT, False),
        ("Implied Share Price", dcf["implied_price"], PRICE_FORMAT, True),
    ]
    for i, (label, value, fmt, bold) in enumerate(bridge_items):
        row = 10 + i
        ws.cell(row=row, column=2, value=label).font = SUBTOTAL_FONT if bold else BODY_FONT
        c = ws.cell(row=row, column=3, value=value)
        c.number_format = fmt
        c.alignment = RIGHT_ALIGN
        c.font = SUBTOTAL_FONT if bold else BODY_FONT
        if bold:
            c.border = THIN_BORDER

    # Highlight implied price
    implied_cell = ws.cell(row=21, column=3)
    implied_cell.fill = GOLD_FILL
    implied_cell.font = HIGHLIGHT_FONT
    ws.cell(row=21, column=2).font = HIGHLIGHT_FONT

    # WACC Calculation (rows 24-35)
    r = 24
    ws.cell(row=r, column=2, value="WACC Calculation").font = HIGHLIGHT_FONT
    ws.cell(row=r, column=2).border = BOTTOM_BORDER

    wacc_items = [
        ("Equity Value (Market Cap, $M)", wacc_calc["equity_value"], CURRENCY_FORMAT),
        ("Debt Value ($M)", wacc_calc["debt_value"], CURRENCY_FORMAT),
        ("Cost of Debt", wacc_calc["cost_of_debt"], PERCENT_FORMAT),
        ("Tax Rate", wacc_calc["tax_rate"], PERCENT_FORMAT),
        ("D / (D+E)", wacc_calc["d_weight"], PERCENT_FORMAT),
        ("After-tax Cost of Debt", wacc_calc["after_tax_cost_of_debt"], PERCENT_FORMAT),
        ("Risk-free Rate", wacc_calc["risk_free_rate"], PERCENT_FORMAT),
        ("Equity Risk Premium", wacc_calc["equity_risk_premium"], PERCENT_FORMAT),
        ("Levered Beta", wacc_calc["beta"], RATIO_FORMAT),
        ("Cost of Equity", wacc_calc["cost_of_equity"], PERCENT_FORMAT),
        ("WACC", wacc_calc["wacc"], PERCENT_FORMAT),
    ]
    for i, (label, value, fmt) in enumerate(wacc_items):
        row = 25 + i
        ws.cell(row=row, column=2, value=label).font = BODY_FONT
        c = ws.cell(row=row, column=3, value=value)
        c.number_format = fmt
        c.alignment = RIGHT_ALIGN
        if label == "WACC":
            c.font = SUBTOTAL_FONT
            c.border = DOUBLE_BOTTOM_BORDER
            ws.cell(row=row, column=2).font = SUBTOTAL_FONT
