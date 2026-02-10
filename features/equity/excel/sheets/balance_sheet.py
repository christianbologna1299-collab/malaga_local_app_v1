"""
Sheet 4: Balance Sheet — Historic + Projected with balance check.
"""

from features.equity.excel.styles import (
    apply_title_row, apply_header_row, apply_freeze_panes, set_column_widths,
    BODY_FONT, SUBTOTAL_FONT, GRAND_TOTAL_FONT, ASSUMPTIONS_FILL,
    GREEN_FILL, RED_FILL,
    CURRENCY_FORMAT, NEGATIVE_CURRENCY, PERCENT_FORMAT,
    LEFT_ALIGN, RIGHT_ALIGN, THIN_BORDER, BOTTOM_BORDER, DOUBLE_BOTTOM_BORDER,
)
from openpyxl.utils import get_column_letter


def build_balance_sheet_sheet(ws, snapshot, outputs):
    """Populate Balance Sheet with historic + projected data and balance check."""
    bs = snapshot["balance_sheet"]
    proj = outputs["projected_balance_sheet"]
    a = outputs["assumptions"]

    h_years = bs["years"]
    p_years = proj["years"]
    all_years = h_years + p_years
    n_cols = len(all_years)
    data_start_col = 3

    widths = {"A": 2, "B": 30}
    for i in range(n_cols):
        widths[get_column_letter(data_start_col + i)] = 16
    set_column_widths(ws, widths)

    # Title
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=data_start_col + n_cols - 1)
    ws["B1"] = "Balance Sheet"
    apply_title_row(ws, 1, data_start_col + n_cols - 1)

    # Column headers
    ws.cell(row=3, column=2, value="$ in millions")
    for i, yr in enumerate(all_years):
        label = f"FY{yr}"
        if yr in p_years:
            label += "E"
        ws.cell(row=3, column=data_start_col + i, value=label).alignment = RIGHT_ALIGN
    apply_header_row(ws, 3, data_start_col + n_cols - 1)
    apply_freeze_panes(ws, "C4")

    def _write_row(row, label, hist_vals, proj_vals, fmt=NEGATIVE_CURRENCY, bold=False, border=None):
        ws.cell(row=row, column=2, value=label).font = SUBTOTAL_FONT if bold else BODY_FONT
        for i, v in enumerate(hist_vals + proj_vals):
            c = ws.cell(row=row, column=data_start_col + i, value=v)
            c.number_format = fmt
            c.alignment = RIGHT_ALIGN
            c.font = SUBTOTAL_FONT if bold else BODY_FONT
            if border:
                c.border = border

    # Current Assets
    r = 5
    _write_row(r, "Cash & Equivalents", bs["cash"], proj["cash"])
    r = 6
    _write_row(r, "Accounts Receivable", bs["ar"], proj["ar"])
    r = 7
    _write_row(r, "Total Current Assets", bs["total_current"], proj["total_current"], bold=True, border=THIN_BORDER)

    # Non-Current Assets
    r = 9
    _write_row(r, "Fixed Assets (Net)", bs["fixed_assets"], proj["fixed_assets"])
    r = 10
    _write_row(r, "Total Non-Current Assets", bs["fixed_assets"], proj["fixed_assets"], bold=True, border=THIN_BORDER)

    # Total Assets
    r = 12
    _write_row(r, "Total Assets", bs["total_assets"], proj["total_assets"], bold=True, border=DOUBLE_BOTTOM_BORDER)
    ws.cell(row=12, column=2).font = GRAND_TOTAL_FONT

    # Current Liabilities
    r = 14
    _write_row(r, "Accounts Payable", bs["ap"], proj["ap"])
    r = 15
    _write_row(r, "Deferred Revenue", bs["deferred_rev"], proj["deferred_rev"])
    r = 16
    _write_row(r, "Total Current Liabilities", bs["total_current_li"], proj["total_current_li"], bold=True, border=THIN_BORDER)

    # Non-Current Liabilities
    r = 18
    _write_row(r, "Long-term Debt", bs["debt"], proj["debt"])
    r = 19
    _write_row(r, "Total Non-Current Liabilities", bs["debt"], proj["debt"], bold=True, border=THIN_BORDER)

    # Total Liabilities
    r = 21
    _write_row(r, "Total Liabilities", bs["total_liabilities"], proj["total_liabilities"], bold=True)

    # Equity
    r = 23
    _write_row(r, "Common Stock", bs["common_stock"], proj["common_stock"])
    r = 24
    _write_row(r, "Retained Earnings", bs["retained_earnings"], proj["retained_earnings"])
    r = 25
    _write_row(r, "Total Equity", bs["total_equity"], proj["total_equity"], bold=True, border=THIN_BORDER)

    # Total L&E
    r = 27
    total_le_hist = [bs["total_liabilities"][i] + bs["total_equity"][i] for i in range(len(h_years))]
    total_le_proj = [proj["total_liabilities"][i] + proj["total_equity"][i] for i in range(len(p_years))]
    _write_row(r, "Total Liabilities & Equity", total_le_hist, total_le_proj, bold=True, border=DOUBLE_BOTTOM_BORDER)
    ws.cell(row=27, column=2).font = GRAND_TOTAL_FONT

    # Balance Check
    r = 29
    ws.cell(row=r, column=2, value="Balance Check").font = SUBTOTAL_FONT
    check_hist = [bs["total_assets"][i] - (bs["total_liabilities"][i] + bs["total_equity"][i]) for i in range(len(h_years))]
    check_proj = [proj["total_assets"][i] - (proj["total_liabilities"][i] + proj["total_equity"][i]) for i in range(len(p_years))]
    for i, v in enumerate(check_hist + check_proj):
        c = ws.cell(row=r, column=data_start_col + i, value=v)
        c.number_format = NEGATIVE_CURRENCY
        c.alignment = RIGHT_ALIGN
        c.font = SUBTOTAL_FONT
        c.fill = GREEN_FILL if v == 0 else RED_FILL

    # Assumptions
    r = 31
    ws.cell(row=r, column=2, value="Assumptions").font = SUBTOTAL_FONT
    ws.cell(row=r, column=2).border = BOTTOM_BORDER
    assumptions_items = [
        ("DSO (days)", a["dso"], "0"),
        ("DPO (days)", a["dpo"], "0"),
        ("CapEx % of Revenue", a["capex_pct_revenue"], PERCENT_FORMAT),
        ("Cost of Debt", a["cost_of_debt"], PERCENT_FORMAT),
    ]
    for i, (label, value, fmt) in enumerate(assumptions_items):
        row = 32 + i
        ws.cell(row=row, column=2, value=label).font = BODY_FONT
        ws.cell(row=row, column=2).fill = ASSUMPTIONS_FILL
        vc = ws.cell(row=row, column=3, value=value)
        vc.number_format = fmt
        vc.alignment = RIGHT_ALIGN
        vc.fill = ASSUMPTIONS_FILL
